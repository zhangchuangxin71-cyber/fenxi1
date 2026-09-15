from __future__ import annotations

"""Ingestion API 容器模块。

该文件集中定义了：
1) 请求/响应模型与状态机；
2) 任务存储（PostgreSQL）；
3) 入库执行器与任务调度服务；
4) FastAPI 路由与应用装配逻辑。
"""

import asyncio
import contextlib
from datetime import datetime, timedelta, timezone
from dataclasses import dataclass, field
from enum import Enum
import hashlib
import hmac
import ipaddress
import json
import logging
import math
import os
from pathlib import Path
import random
import re
import secrets
import shutil
import sys
import tempfile
import time
from typing import Any, Awaitable, Callable, Literal, Protocol
from urllib.parse import unquote, urlparse
import uuid

import requests
from pydantic import BaseModel, ConfigDict, Field, HttpUrl, PrivateAttr, field_validator, model_validator

_REPO_ROOT = Path(__file__).resolve().parents[2]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

from app.config.env import bootstrap_runtime_env
from app.ingestion.agent_adapter import DocumentAssistantAdapter
from app.parser.conversion_core.config import RuntimeConfig
from app.parser.mineru_adapter import parse_document_with_mineru_async
from app.storage.postgres_store import is_complete_raw_mineru
from app.storage.errors import DuplicateDocumentError

bootstrap_runtime_env(_REPO_ROOT)

SUPPORTED_API_FILE_TYPES = Literal[
    "pdf",
    "doc",
    "docx",
    "txt",
    "md",
    "markdown",
    "html",
    "xlsx",
    "pptx",
]


def utc_now_iso() -> str:
    """返回 UTC ISO8601 时间字符串，统一使用 `Z` 结尾。"""
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


class TaskStatus(str, Enum):
    """任务生命周期状态枚举。"""
    QUEUED = "queued"
    PARSING = "parsing"
    EXTRACTING = "extracting"
    STRUCTURING = "structuring"
    STORING = "storing"
    COMPLETED = "completed"
    FAILED = "failed"


IN_PROGRESS_STATUSES = {
    TaskStatus.QUEUED,
    TaskStatus.PARSING,
    TaskStatus.EXTRACTING,
    TaskStatus.STRUCTURING,
    TaskStatus.STORING,
}


class IngestConfig(BaseModel):
    """入库执行参数配置。"""
    model_config = ConfigDict(extra="forbid")

    # 合约优先字段（A1 / A3）。
    enable_ocr: bool = Field(default=True, description="是否启用 OCR。")
    language: Literal["zh", "en", "auto"] = Field(default="zh", description="文档主语言。")
    extract_tables: bool = Field(default=True, description="是否提取表格。")
    extract_images: bool = Field(default=False, description="是否提取图片描述。")
    max_tree_depth: Literal[1, 2, 3] = Field(default=3, description="PageIndex 树最大深度。")
    chunk_overlap: int = Field(default=50, description="相邻叶子节点文本重叠字符数。")
    node_max_tokens: int = Field(default=512, description="叶子节点最大 Token 数。")
    summary_enabled: bool = Field(default=True, description="是否开启摘要生成。")
    table_parse_mode: Literal["off", "auto", "rule", "model"] = Field(default="auto", description="表格解析模式。")

    # 兼容旧调用方的字段。
    chunk_overlap_ratio: float | None = Field(default=None, description="分块重叠比例（兼容字段，优先级低于 chunk_overlap）。")

    @field_validator("node_max_tokens")
    @classmethod
    def validate_node_max_tokens(cls, v: int) -> int:
        """校验节点最大 token 配置是否合法。"""
        if v < 64:
            raise ValueError("node_max_tokens must be >= 64")
        return v

    @field_validator("chunk_overlap_ratio")
    @classmethod
    def validate_overlap(cls, v: float | None) -> float | None:
        """校验分块重叠比例是否在允许范围。"""
        if v is None:
            return None
        if not (0.1 <= float(v) <= 0.5):
            raise ValueError("chunk_overlap_ratio must be within [0.1, 0.5]")
        return float(v)

    @field_validator("chunk_overlap")
    @classmethod
    def validate_chunk_overlap(cls, v: int) -> int:
        """校验重叠字符数配置。"""
        if int(v) < 0:
            raise ValueError("chunk_overlap must be >= 0")
        return int(v)

    def resolved_overlap_ratio(self) -> float:
        """将 chunk_overlap / chunk_overlap_ratio 统一映射到 ratio（0.1~0.5）。"""
        if self.chunk_overlap_ratio is not None:
            return max(0.1, min(0.5, float(self.chunk_overlap_ratio)))
        # 经验映射：50 个字符约等于 0.1 的重叠比例，最大不超过 0.5。
        ratio = float(self.chunk_overlap) / 500.0
        return max(0.1, min(0.5, ratio))

    def resolved_table_parse_mode(self) -> str:
        """统一表格解析模式；extract_tables=false 时强制 off。"""
        if not bool(self.extract_tables):
            return "off"
        mode = str(self.table_parse_mode or "auto").strip().lower()
        if mode in {"rule", "model"}:
            # 当前入库服务只稳定支持 off / auto。
            return "auto"
        return "off" if mode == "off" else "auto"


class IngestDocument(BaseModel):
    """批量提交时的单文档描述。"""
    model_config = ConfigDict(extra="forbid")

    _local_path: str | None = PrivateAttr(default=None)

    oss_key: str = Field(min_length=1, description="OSS 对象 Key。")
    file_name: str = Field(min_length=1, description="文件名（用于展示与回退检索）。")
    file_type: SUPPORTED_API_FILE_TYPES = Field(description="文件类型。")

    @field_validator("oss_key")
    @classmethod
    def validate_oss_key(cls, v: str) -> str:
        """校验 OSS key 字段非空且格式可用。"""
        value = str(v or "").strip().lstrip("/")
        if not value:
            raise ValueError("oss_key is required")
        return value


class IngestRequest(BaseModel):
    """标准入库请求：支持单文档字段或 documents 批量字段。"""
    model_config = ConfigDict(extra="forbid")

    _doc_id: str = PrivateAttr(default_factory=lambda: str(uuid.uuid4()))
    _doc_ids: list[str] = PrivateAttr(default_factory=list)
    _submit_temp_dir: str | None = PrivateAttr(default=None)
    _single_local_path: str | None = PrivateAttr(default=None)
    _raw_mineru_repair_doc_ids: set[str] = PrivateAttr(default_factory=set)
    _raw_mineru_repair_id: str | None = PrivateAttr(default=None)

    user_id: str = Field(min_length=1, description="用户 ID，用于标识多用户平台中的文档归属。")
    kb_id: str = Field(min_length=1, description="知识库 ID。")
    oss_key: str | None = Field(default=None, description="单文件模式：OSS 对象 Key。")
    file_name: str | None = Field(default=None, description="单文件模式：文件名。")
    file_type: SUPPORTED_API_FILE_TYPES | None = Field(default=None, description="单文件模式：文件类型。")
    documents: list[IngestDocument] = Field(default_factory=list, description="批量模式：文档列表。")
    callback_url: HttpUrl | None = Field(default=None, description="任务完成/失败后的回调地址。")
    idempotency_key: str | None = Field(default=None, description="幂等键；同一 kb_id 下重复提交可命中历史任务。")
    is_temp: bool = Field(default=False, description="是否为临时文档任务。")
    session_id: str | None = Field(default=None, description="临时文档关联会话 ID。")
    config: IngestConfig = Field(default_factory=IngestConfig, description="入库算法参数配置。")

    @property
    def doc_id(self) -> str:
        """返回内部任务文档 ID；对外接口不再要求调用方传入。"""
        return self._doc_id

    @property
    def doc_ids(self) -> list[str]:
        """返回本次请求每个文档对应的 UUID。"""
        return list(self._doc_ids or [self._doc_id])

    @property
    def raw_mineru_repair_doc_ids(self) -> set[str]:
        """Return document IDs scheduled for an in-place MinerU raw repair."""
        return set(self._raw_mineru_repair_doc_ids)
    @property
    def raw_mineru_repair_id(self) -> str | None:
        """Return the active on-demand repair claim ID, when present."""
        return self._raw_mineru_repair_id


    @model_validator(mode="after")
    def validate_payload_shape(self) -> "IngestRequest":
        """校验单文件模式与批量模式的请求结构互斥且完整。"""
        has_single = bool(
            str(self.oss_key or "").strip()
            or str(self.file_name or "").strip()
            or self.file_type is not None
        )
        has_batch = bool(self.documents)
        if has_single and has_batch:
            raise ValueError("单文件字段和 documents 批量字段只能二选一")
        if not has_single and not has_batch:
            raise ValueError("必须提供 oss_key，并提供 file_name/file_type；或使用 documents")
        if has_single:
            has_oss_key = bool(str(self.oss_key or "").strip())
            if not has_oss_key:
                raise ValueError("oss_key is required")
            if not str(self.file_name or "").strip():
                raise ValueError("file_name is required")
            if self.file_type is None:
                raise ValueError("file_type is required")
        if has_batch and len(self.documents) < 1:
            raise ValueError("documents must not be empty")
        return self

    @field_validator("oss_key")
    @classmethod
    def validate_oss_key(cls, v: str | None) -> str | None:
        """校验 OSS key 字段非空且格式可用。"""
        if v is None:
            return None
        value = str(v or "").strip().lstrip("/")
        if not value:
            raise ValueError("oss_key is required")
        return value

    def iter_documents(self) -> list[IngestDocument]:
        """统一返回待处理文档列表（兼容单文件与批量）。"""
        if self.documents:
            return list(self.documents)
        doc = IngestDocument(
            oss_key=str(self.oss_key or ""),
            file_name=str(self.file_name or ""),
            file_type=self.file_type or "md",
        )
        doc._local_path = self._single_local_path
        return [doc]

    @property
    def document_count(self) -> int:
        """返回本次请求的文档数量。"""
        return len(self.iter_documents())


class TempIngestRequest(BaseModel):
    """临时文档入库请求（高优先级队列）。"""
    model_config = ConfigDict(extra="forbid")

    user_id: str = Field(min_length=1, description="用户 ID，用于标识多用户平台中的临时文档归属。")
    session_id: str = Field(min_length=1, description="会话 ID，用于区分临时会话范围。")
    kb_id: str = Field(min_length=1, description="知识库 ID。")
    oss_key: str = Field(min_length=1, description="OSS 对象 Key。")
    file_name: str = Field(min_length=1, description="文件名。")
    file_type: SUPPORTED_API_FILE_TYPES = Field(description="文件类型。")
    callback_url: HttpUrl | None = Field(default=None, description="任务完成/失败后的回调地址。")
    idempotency_key: str | None = Field(default=None, description="幂等键。")
    config: IngestConfig = Field(default_factory=IngestConfig, description="入库算法参数配置。")

    @field_validator("oss_key")
    @classmethod
    def validate_oss_key(cls, v: str) -> str:
        """校验 OSS key 字段非空且格式可用。"""
        value = str(v or "").strip().lstrip("/")
        if not value:
            raise ValueError("oss_key is required")
        return value


class IngestSubmitData(BaseModel):
    """提交接口返回的数据体。"""
    task_id: str = Field(description="任务 ID，可用于查询任务状态。")
    doc_id: str = Field(description="内部任务文档 ID。")
    doc_ids: list[str] = Field(default_factory=list, description="批量场景下每个文档对应的 UUID。")
    user_id: str = Field(description="用户 ID。")
    status: Literal["queued", "completed"] = Field(default="queued", description="提交结果状态。")
    estimated_seconds: int = Field(description="预估剩余秒数（粗略值）。")
    is_duplicate: bool = Field(default=False, description="是否命中幂等（true 表示返回历史任务）。")
    document_count: int = Field(default=1, description="本次提交的文档数量。")


class TaskStatusData(BaseModel):
    """任务状态详情（用于轮询查询与回调）。"""
    doc_id: str = Field(description="文档 ID。")
    user_id: str = Field(default="system", description="用户 ID。")
    kb_id: str = Field(description="知识库 ID。")
    task_id: str = Field(description="任务 ID。")
    status: TaskStatus = Field(description="任务状态。")
    progress: int = Field(default=0, description="任务进度（0~100）。")
    current_step: str = Field(default="", description="当前步骤说明。")
    total_pages: int | None = Field(default=None, description="总页数（可为空）。")
    processed_pages: int = Field(default=0, description="已处理页数。")
    tree_node_count: int = Field(default=0, description="已生成结构节点数量。")
    error_message: str | None = Field(default=None, description="失败时的错误信息。")
    error_code: str | None = Field(default=None, description="失败时的错误码。")
    retryable: bool = Field(default=False, description="失败是否可重试。")
    started_at: str | None = Field(default=None, description="开始时间（UTC ISO8601）。")
    updated_at: str = Field(description="最近更新时间（UTC ISO8601）。")
    completed_at: str | None = Field(default=None, description="完成时间（UTC ISO8601）。")
    retry_count: int = Field(default=0, description="重试次数。")
    eta_seconds: int | None = Field(default=None, description="预计剩余秒数。")
    persisted_doc_id: str | None = Field(default=None, description="落库后的主文档 ID。")
    persisted_doc_ids: list[str] = Field(default_factory=list, description="批量场景下落库后的文档 ID 列表。")
    total_documents: int = Field(default=1, description="任务总文档数。")
    processed_documents: int = Field(default=0, description="已处理文档数。")
    callback_url: str | None = Field(default=None, description="回调地址。")
    idempotency_key: str | None = Field(default=None, description="幂等键。")
    is_temp: bool = Field(default=False, description="是否为临时任务。")
    session_id: str | None = Field(default=None, description="临时任务关联会话 ID。")


class ApiResponse(BaseModel):
    """统一 API 响应包装。"""
    code: int = Field(description="业务状态码。")
    message: str = Field(description="提示信息。")
    data: Any = Field(default=None, description="返回数据体。")

class RawMineruRepairRequest(BaseModel):
    """Internal request to repair MinerU raw output in place."""

    user_id: str = Field(min_length=1)
    kb_id: str = Field(min_length=1)
    doc_id: str = Field(min_length=1)

def _public_status_payload(task: TaskStatusData) -> dict[str, Any]:
    """按 A2 契约返回状态字段（避免泄露内部扩展字段）。"""
    return {
        "doc_id": task.doc_id,
        "user_id": task.user_id,
        "task_id": task.task_id,
        "status": task.status.value if isinstance(task.status, TaskStatus) else str(task.status),
        "progress": int(task.progress),
        "current_step": task.current_step,
        "total_pages": task.total_pages,
        "processed_pages": int(task.processed_pages),
        "tree_node_count": int(task.tree_node_count),
        "error_message": task.error_message,
        "started_at": task.started_at,
        "completed_at": task.completed_at,
    }


def _env_bool(name: str, default: bool = False) -> bool:
    """从环境变量解析布尔值。"""
    raw = str(os.getenv(name, "") or "").strip().lower()
    if not raw:
        return default
    return raw in {"1", "true", "yes", "on"}


def _split_csv(raw: str) -> list[str]:
    """将逗号分隔字符串转换为小写列表。"""
    return [x.strip().lower() for x in str(raw or "").split(",") if x.strip()]


def _env_str(name: str, default: str = "") -> str:
    """读取字符串环境变量，自动去除首尾空白。"""
    return str(os.getenv(name, default) or default).strip()


def _env_path_list(name: str) -> list[Path]:
    """读取 pathsep 分隔的路径列表环境变量。"""
    raw = _env_str(name, "")
    if not raw:
        return []
    paths: list[Path] = []
    for token in raw.split(os.pathsep):
        value = str(token or "").strip()
        if value:
            paths.append(Path(value).expanduser())
    return paths


def _env_int(name: str, default: int) -> int:
    """读取整数环境变量；非法值时回退默认值。"""
    raw = str(os.getenv(name, "") or "").strip()
    if not raw:
        return int(default)
    try:
        return int(raw)
    except Exception:
        logging.getLogger(__name__).warning("invalid %s=%r; fallback=%s", name, raw, default)
        return int(default)


def _normalize_http_path(path: str, default: str) -> str:
    """规范化 HTTP 路径，保证以 `/` 开头。"""
    text = str(path or "").strip()
    if not text:
        text = default
    return text if text.startswith("/") else "/" + text


_IDENTIFIER_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def _normalize_identifier(value: str, *, default: str, label: str) -> str:
    """规范化数据库标识符，不合法时回退默认值。"""
    text = str(value or "").strip()
    if text and _IDENTIFIER_RE.match(text):
        return text
    logging.getLogger(__name__).warning("invalid %s=%r; fallback=%s", label, value, default)
    return default


def _normalize_storage_doc_type(file_type: str | None) -> str:
    """Normalize legacy Office document uploads to the existing stored docx type."""
    normalized = str(file_type or "").strip().lower().lstrip(".")
    return "docx" if normalized == "doc" else normalized


API_PREFIX = _normalize_http_path(_env_str("INGEST_API_PREFIX", "/ingestion/v1"), "/ingestion/v1")
HEALTH_PATH = _normalize_http_path(_env_str("INGEST_HEALTH_PATH", "/healthz"), "/healthz")
WORKSPACE_DIR = _env_str("INGEST_WORKSPACE_DIR", "")
ALLOWED_INPUT_SCHEMES = set(_split_csv(_env_str("INGEST_ALLOWED_INPUT_SCHEMES", "http,https,file,local")))
if not ALLOWED_INPUT_SCHEMES:
    ALLOWED_INPUT_SCHEMES = {"http", "https", "file", "local"}
CALLBACK_ALLOWED_SCHEMES = set(_split_csv(_env_str("INGEST_CALLBACK_ALLOWED_SCHEMES", "http,https")))
if not CALLBACK_ALLOWED_SCHEMES:
    CALLBACK_ALLOWED_SCHEMES = {"http", "https"}
DEFAULT_SEARCH_ROOTS_ENV = _env_path_list("INGEST_DEFAULT_SEARCH_ROOTS")
CLEANUP_SCHEMA = _normalize_identifier(_env_str("INGEST_CLEANUP_SCHEMA", "public"), default="public", label="INGEST_CLEANUP_SCHEMA")
CLEANUP_TABLES_RAW = _split_csv(_env_str("INGEST_CLEANUP_TABLES", "doc_nodes,doc_pages,documents"))
CLEANUP_TABLES: list[str] = []
for _name in CLEANUP_TABLES_RAW:
    _candidate = str(_name or "").strip()
    if _candidate and _IDENTIFIER_RE.match(_candidate):
        CLEANUP_TABLES.append(_candidate)
    elif _candidate:
        logging.getLogger(__name__).warning("invalid INGEST_CLEANUP_TABLES item=%r; skipped", _candidate)
if not CLEANUP_TABLES:
    CLEANUP_TABLES = ["doc_nodes", "doc_pages", "documents"]
TREE_NODES_TABLE = _normalize_identifier(
    _env_str("INGEST_TREE_NODES_TABLE", CLEANUP_TABLES[0]),
    default=CLEANUP_TABLES[0],
    label="INGEST_TREE_NODES_TABLE",
)


@dataclass
class CallbackDLQItem:
    """回调死信队列条目：用于记录回调最终失败的信息。"""
    task_id: str
    doc_id: str
    callback_url: str
    last_http_status: int | None
    last_error: str
    attempt_count: int
    next_action: str = "manual_replay"
    created_at: float = field(default_factory=time.time)
    updated_at: float = field(default_factory=time.time)


class CallbackDispatcher:
    """任务状态回调分发器（含白名单校验、签名与重试）。"""
    def __init__(self) -> None:
        """初始化对象并准备运行所需的配置与状态。"""
        self.enabled = _env_bool("INGEST_CALLBACK_ENABLED", True)
        self.whitelist = _split_csv(os.getenv("INGEST_CALLBACK_WHITELIST", ""))
        self.allowed_schemes = set(CALLBACK_ALLOWED_SCHEMES)
        self.secret = str(os.getenv("INGEST_CALLBACK_SECRET", "") or "").strip()
        self.connect_timeout = float(os.getenv("INGEST_CALLBACK_CONNECT_TIMEOUT", "3"))
        self.read_timeout = float(os.getenv("INGEST_CALLBACK_READ_TIMEOUT", "5"))
        self.max_retries = int(os.getenv("INGEST_CALLBACK_MAX_RETRIES", "5"))
        self.base_backoff = float(os.getenv("INGEST_CALLBACK_BASE_BACKOFF", "2"))
        self.jitter_ratio = float(os.getenv("INGEST_CALLBACK_JITTER_RATIO", "0.2"))
        self.dlq: list[CallbackDLQItem] = []

    def _is_private_or_loopback(self, host: str) -> bool:
        """判断主机是否为私网、环回或保留地址。"""
        try:
            ip = ipaddress.ip_address(host)
            return ip.is_loopback or ip.is_private or ip.is_link_local or ip.is_reserved
        except Exception:
            return False

    def _is_allowed(self, callback_url: str) -> bool:
        """按白名单与安全规则校验回调地址是否允许访问。"""
        parsed = urlparse(callback_url)
        if parsed.scheme not in self.allowed_schemes:
            return False
        host = str(parsed.hostname or "").strip().lower()
        if not host:
            return False
        if self._is_private_or_loopback(host):
            return False
        # When whitelist is empty, allow public callback domains by default.
        if not self.whitelist:
            return True
        return any(host == allowed or host.endswith("." + allowed) for allowed in self.whitelist)

    def _sign_headers(self, method: str, path: str, body: dict[str, Any]) -> dict[str, str]:
        """为回调请求构造签名请求头。"""
        timestamp = str(int(time.time()))
        nonce = secrets.token_hex(12)
        body_bytes = json.dumps(body, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        body_sha256 = hashlib.sha256(body_bytes).hexdigest()
        signing = f"{timestamp}\n{nonce}\n{method.upper()}\n{path}\n{body_sha256}"
        signature = hmac.new(self.secret.encode("utf-8"), signing.encode("utf-8"), hashlib.sha256).hexdigest()
        return {
            "X-Callback-Timestamp": timestamp,
            "X-Callback-Nonce": nonce,
            "X-Callback-Signature": signature,
            "Content-Type": "application/json",
        }

    @staticmethod
    def _retryable_status(status: int) -> bool:
        """判断 HTTP 状态码是否可重试。"""
        return status == 429 or status >= 500

    def _compute_sleep(self, attempt: int) -> float:
        """计算重试退避等待时长（含抖动）。"""
        base = self.base_backoff * (2 ** max(0, attempt - 1))
        jitter = base * self.jitter_ratio * random.random()
        return base + jitter

    def _build_payload(self, task: TaskStatusData) -> dict[str, Any]:
        """构造回调请求体。"""
        return {
            "task_id": task.task_id,
            "doc_id": task.doc_id,
            "user_id": task.user_id,
            "kb_id": task.kb_id,
            "status": task.status.value,
            "progress": task.progress,
            "current_step": task.current_step,
            "error_message": task.error_message,
            "error_code": task.error_code,
            "retryable": task.retryable,
            "started_at": task.started_at,
            "updated_at": task.updated_at,
            "completed_at": task.completed_at,
        }

    async def dispatch(self, callback_url: str | None, task: TaskStatusData) -> None:
        """发送任务回调；失败时按策略重试并落入 DLQ。"""
        url = str(callback_url or "").strip()
        if not self.enabled or not url:
            return
        if not self._is_allowed(url):
            self.dlq.append(
                CallbackDLQItem(
                    task_id=task.task_id,
                    doc_id=task.doc_id,
                    callback_url=url,
                    last_http_status=None,
                    last_error="callback_url not allowed",
                    attempt_count=0,
                    next_action="abandon",
                )
            )
            return

        parsed = urlparse(url)
        path = parsed.path or "/"
        payload = self._build_payload(task)
        headers = self._sign_headers("POST", path, payload) if self.secret else {"Content-Type": "application/json"}

        last_status: int | None = None
        last_error = ""
        for attempt in range(0, self.max_retries + 1):
            try:
                resp = await self._post(url, payload, headers)
                last_status = int(resp.status_code)
                if 200 <= resp.status_code < 300:
                    return
                if not self._retryable_status(resp.status_code):
                    last_error = f"non-retryable status={resp.status_code}"
                    break
                last_error = f"retryable status={resp.status_code}"
            except requests.RequestException as exc:
                last_error = str(exc)
            if attempt < self.max_retries:
                await self._sleep(self._compute_sleep(attempt + 1))

        self.dlq.append(
            CallbackDLQItem(
                task_id=task.task_id,
                doc_id=task.doc_id,
                callback_url=url,
                last_http_status=last_status,
                last_error=last_error or "callback failed",
                attempt_count=self.max_retries + 1,
            )
        )

    async def _post(self, url: str, payload: dict[str, Any], headers: dict[str, str]) -> requests.Response:
        """通过线程池发送 HTTP POST 请求。"""
        return await asyncio.to_thread(
            requests.post,
            url,
            json=payload,
            headers=headers,
            timeout=(self.connect_timeout, self.read_timeout),
        )

    async def _sleep(self, seconds: float) -> None:
        """异步等待指定秒数。"""
        await asyncio.sleep(seconds)




try:
    import psycopg2
    from psycopg2 import sql as psycopg2_sql
    from psycopg2.extras import RealDictCursor
except Exception:  # pragma: no cover - optional dependency
    psycopg2 = None  # type: ignore[assignment]
    psycopg2_sql = None  # type: ignore[assignment]
    RealDictCursor = None  # type: ignore[assignment]



CREATE_TASKS_SQL = """
-- 入库任务主表：记录每个 ingestion 任务的生命周期状态、进度与结果。
CREATE TABLE IF NOT EXISTS ingestion_tasks (
    -- 任务基础标识
    task_id TEXT PRIMARY KEY,
    doc_id TEXT NOT NULL,
    user_id TEXT NOT NULL DEFAULT 'system',
    kb_id TEXT NOT NULL,

    -- 任务执行状态
    status TEXT NOT NULL,
    progress INTEGER NOT NULL DEFAULT 0,
    current_step TEXT NOT NULL DEFAULT '',

    -- 文档处理统计
    total_pages INTEGER NULL,
    processed_pages INTEGER NOT NULL DEFAULT 0,
    tree_node_count INTEGER NOT NULL DEFAULT 0,

    -- 错误与重试信息
    error_message TEXT NULL,
    error_code TEXT NULL,
    retryable BOOLEAN NOT NULL DEFAULT FALSE,
    retry_count INTEGER NOT NULL DEFAULT 0,

    -- 时间轴
    started_at TIMESTAMPTZ NULL,
    updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
    completed_at TIMESTAMPTZ NULL,

    -- 估时与持久化结果
    eta_seconds INTEGER NULL,
    persisted_doc_id TEXT NULL,
    persisted_doc_ids JSONB NOT NULL DEFAULT '[]'::jsonb,

    -- 批量任务统计
    total_documents INTEGER NOT NULL DEFAULT 1,
    processed_documents INTEGER NOT NULL DEFAULT 0,

    -- 回调与幂等信息
    callback_url TEXT NULL,
    idempotency_key TEXT NULL,

    -- 临时任务上下文
    is_temp BOOLEAN NOT NULL DEFAULT FALSE,
    session_id TEXT NULL,

    -- 记录创建时间
    created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
);

-- 常用查询索引（按 doc_id、业务键、更新时间、幂等键）
ALTER TABLE ingestion_tasks ADD COLUMN IF NOT EXISTS user_id TEXT NOT NULL DEFAULT 'system';
CREATE INDEX IF NOT EXISTS idx_ingestion_tasks_doc_id ON ingestion_tasks(doc_id);
CREATE INDEX IF NOT EXISTS idx_ingestion_tasks_user_id ON ingestion_tasks(user_id);
CREATE INDEX IF NOT EXISTS idx_ingestion_tasks_doc_kb ON ingestion_tasks(doc_id, kb_id);
CREATE INDEX IF NOT EXISTS idx_ingestion_tasks_updated_at ON ingestion_tasks(updated_at DESC);
CREATE INDEX IF NOT EXISTS idx_ingestion_tasks_kb_idem_updated ON ingestion_tasks(kb_id, idempotency_key, updated_at DESC);

-- 向后兼容：旧版本表结构缺失列时自动补齐
ALTER TABLE ingestion_tasks ADD COLUMN IF NOT EXISTS callback_url TEXT NULL;
ALTER TABLE ingestion_tasks ADD COLUMN IF NOT EXISTS user_id TEXT NOT NULL DEFAULT 'system';
ALTER TABLE ingestion_tasks ADD COLUMN IF NOT EXISTS is_temp BOOLEAN NOT NULL DEFAULT FALSE;
ALTER TABLE ingestion_tasks ADD COLUMN IF NOT EXISTS session_id TEXT NULL;
ALTER TABLE ingestion_tasks ADD COLUMN IF NOT EXISTS persisted_doc_id TEXT NULL;
ALTER TABLE ingestion_tasks ADD COLUMN IF NOT EXISTS persisted_doc_ids JSONB NOT NULL DEFAULT '[]'::jsonb;
ALTER TABLE ingestion_tasks ADD COLUMN IF NOT EXISTS total_documents INTEGER NOT NULL DEFAULT 1;
ALTER TABLE ingestion_tasks ADD COLUMN IF NOT EXISTS processed_documents INTEGER NOT NULL DEFAULT 0;
ALTER TABLE ingestion_tasks ADD COLUMN IF NOT EXISTS idempotency_key TEXT NULL;
"""


CREATE_EVENTS_SQL = """
CREATE TABLE IF NOT EXISTS ingestion_task_events (
    id BIGSERIAL PRIMARY KEY,
    task_id TEXT NOT NULL,
    status TEXT NOT NULL,
    event_payload JSONB NOT NULL DEFAULT '{}'::jsonb,
    created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
);

CREATE INDEX IF NOT EXISTS idx_ingestion_task_events_task_id ON ingestion_task_events(task_id, created_at DESC);
"""


@dataclass
class _PgConfig:
    """PostgreSQL 连接配置对象。"""
    dsn: str


def resolve_postgres_dsn() -> str:
    """读取并返回 PostgreSQL DSN。"""
    return str(os.getenv("POSTGRES_DSN", "")).strip()


class PostgresTaskStore:
    """基于 PostgreSQL 的任务状态持久化存储。"""
    def __init__(self, dsn: str) -> None:
        """初始化对象并准备运行所需的配置与状态。"""
        if not dsn:
            raise ValueError("PostgreSQL DSN is required for PostgresTaskStore.")
        if psycopg2 is None:
            raise RuntimeError("psycopg2 is required for PostgresTaskStore")
        self._cfg = _PgConfig(dsn=dsn)
        self._init_lock = asyncio.Lock()
        self._initialized = False
        self._event_progress_step = max(1, int(os.getenv("INGEST_EVENT_PROGRESS_STEP", "5")))
        self._event_min_interval_seconds = max(0.0, float(os.getenv("INGEST_EVENT_MIN_INTERVAL_SECONDS", "1")))

    def _connect(self):
        """创建并返回 PostgreSQL 连接。"""
        return psycopg2.connect(self._cfg.dsn)

    def get_binding_doc_name(self, doc_id: str, user_id: str, kb_id: str) -> str | None:
        """Return an existing document name in the requested binding scope."""
        sql = """
        SELECT d.doc_name
        FROM document_bindings b
        JOIN documents d ON d.doc_id = b.doc_id
        WHERE b.doc_id = %s AND b.user_id = %s AND b.kb_id = %s
        LIMIT 1
        """
        with self._connect() as conn:
            with conn.cursor() as cur:
                cur.execute(sql, (doc_id, user_id, kb_id))
                row = cur.fetchone()
        return (str(row[0] or "") or None) if row else None

    def has_complete_raw_mineru(self, doc_id: str) -> bool:
        """Check whether the shared document already has complete MinerU raw data."""
        sql = """
        SELECT raw -> 'raw_mineru'
        FROM documents
        WHERE doc_id = %s
        LIMIT 1
        """
        with self._connect() as conn:
            with conn.cursor() as cur:
                cur.execute(sql, (doc_id,))
                row = cur.fetchone()
        if not row:
            return False
        raw_mineru = row[0]
        if isinstance(raw_mineru, str):
            try:
                raw_mineru = json.loads(raw_mineru)
            except json.JSONDecodeError:
                return False
        return is_complete_raw_mineru(raw_mineru)

    @staticmethod
    def _parse_repair_utc(value: Any) -> datetime | None:
        raw = str(value or "").strip()
        if not raw:
            return None
        try:
            return datetime.fromisoformat(raw.replace("Z", "+00:00")).astimezone(timezone.utc)
        except (TypeError, ValueError):
            return None

    def claim_raw_mineru_repair(
        self,
        *,
        doc_id: str,
        user_id: str,
        kb_id: str,
        repair_id: str,
        lease_seconds: int = 1200,
        retry_cooldown_seconds: int = 60,
        max_attempts: int = 3,
        require_source: bool = True,
    ) -> dict[str, Any]:
        """Atomically claim an on-demand MinerU-only repair."""
        sql = """
        SELECT d.doc_id, d.doc_name, d.doc_type, d.file_oss_key,
               d.raw -> 'raw_mineru_repair', d.raw -> 'raw_mineru'
        FROM document_bindings b
        JOIN documents d ON d.doc_id = b.doc_id
        WHERE b.doc_id = %s AND b.user_id = %s AND b.kb_id = %s
        LIMIT 1
        FOR UPDATE OF d
        """
        now = datetime.now(timezone.utc)
        with self._connect() as conn:
            with conn.cursor() as cur:
                cur.execute(sql, (doc_id, user_id, kb_id))
                row = cur.fetchone()
                if not row:
                    return {"outcome": "not_found"}

                repair = row[4] if isinstance(row[4], dict) else {}
                raw_mineru = row[5]
                if isinstance(raw_mineru, str):
                    with contextlib.suppress(json.JSONDecodeError):
                        raw_mineru = json.loads(raw_mineru)
                if is_complete_raw_mineru(raw_mineru):
                    return {"outcome": "complete", "repair": repair}

                source_key = str(row[3] or "").strip()
                if require_source and not source_key:
                    return {"outcome": "missing_source", "repair": repair}

                status = str(repair.get("status") or "")
                lease_expires_at = self._parse_repair_utc(repair.get("lease_expires_at"))
                if status in {"queued", "processing"} and lease_expires_at and lease_expires_at > now:
                    return {"outcome": "active", "repair": repair}

                attempts = max(0, int(repair.get("attempt_count") or 0))
                if status == "failed" and not bool(repair.get("retryable")):
                    return {"outcome": "terminal_failed", "repair": repair}
                if status == "failed" and attempts >= max(1, int(max_attempts)):
                    repair = {
                        **repair,
                        "retryable": False,
                        "error_code": "RAW_MINERU_REPAIR_ATTEMPTS_EXHAUSTED",
                        "error_message": "MinerU repair retry limit reached",
                        "updated_at": utc_now_iso(),
                        "lease_expires_at": None,
                        "next_retry_at": None,
                    }
                    cur.execute(
                        """
                        UPDATE documents
                        SET raw = jsonb_set(
                          raw, '{raw_mineru_repair}', %s::jsonb, true
                        )
                        WHERE doc_id = %s
                          AND raw -> 'raw_mineru_repair' ->> 'repair_id' = %s
                        """,
                        (
                            json.dumps(repair, ensure_ascii=False),
                            doc_id,
                            str(repair.get("repair_id") or ""),
                        ),
                    )
                    conn.commit()
                    return {"outcome": "terminal_failed", "repair": repair}
                next_retry_at = self._parse_repair_utc(repair.get("next_retry_at"))
                if status == "failed" and next_retry_at and next_retry_at > now:
                    return {"outcome": "cooldown", "repair": repair}

                timestamp = utc_now_iso()
                repair = {
                    "repair_id": repair_id,
                    "status": "queued",
                    "attempt_count": attempts + 1,
                    "created_at": repair.get("created_at") or timestamp,
                    "updated_at": timestamp,
                    "lease_expires_at": (
                        now + timedelta(seconds=max(1, int(lease_seconds)))
                    ).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
                    "next_retry_at": (
                        now + timedelta(seconds=max(0, int(retry_cooldown_seconds)))
                    ).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
                    "retryable": True,
                }
                cur.execute(
                    """
                    UPDATE documents
                    SET raw = jsonb_set(
                      CASE WHEN jsonb_typeof(raw) = 'object' THEN raw ELSE '{}'::jsonb END,
                      '{raw_mineru_repair}', %s::jsonb, true
                    )
                    WHERE doc_id = %s
                    """,
                    (json.dumps(repair, ensure_ascii=False), doc_id),
                )
            conn.commit()
        return {
            "outcome": "claimed",
            "repair": repair,
            "document": {
                "doc_id": str(row[0]),
                "doc_name": str(row[1] or ""),
                "doc_type": str(row[2] or "").lower(),
                "file_oss_key": source_key,
            },
        }

    def get_raw_mineru_repair(
        self,
        *,
        doc_id: str,
        user_id: str,
        kb_id: str,
    ) -> dict[str, Any] | None:
        """Read persisted repair state without claiming or retrying it."""
        sql = """
        SELECT d.raw -> 'raw_mineru_repair', d.raw -> 'raw_mineru'
        FROM document_bindings b
        JOIN documents d ON d.doc_id = b.doc_id
        WHERE b.doc_id = %s AND b.user_id = %s AND b.kb_id = %s
        LIMIT 1
        """
        with self._connect() as conn:
            with conn.cursor() as cur:
                cur.execute(sql, (doc_id, user_id, kb_id))
                row = cur.fetchone()
        if not row:
            return None
        repair = dict(row[0]) if isinstance(row[0], dict) else {}
        raw_mineru = row[1]
        if isinstance(raw_mineru, str):
            with contextlib.suppress(json.JSONDecodeError):
                raw_mineru = json.loads(raw_mineru)
        if is_complete_raw_mineru(raw_mineru):
            repair.update({"status": "completed", "retryable": False})
        elif not repair:
            repair = {"status": "not_started", "retryable": True}
        return repair

    def update_raw_mineru_repair(
        self,
        *,
        doc_id: str,
        repair_id: str,
        updates: dict[str, Any],
    ) -> None:
        """Patch repair metadata only when the caller still owns the active repair."""
        payload = dict(updates)
        payload["updated_at"] = utc_now_iso()
        sql = """
        UPDATE documents
        SET raw = jsonb_set(
          CASE WHEN jsonb_typeof(raw) = 'object' THEN raw ELSE '{}'::jsonb END,
          '{raw_mineru_repair}',
          (
            CASE
              WHEN jsonb_typeof(raw -> 'raw_mineru_repair') = 'object'
                THEN raw -> 'raw_mineru_repair'
              ELSE '{}'::jsonb
            END
            || %s::jsonb
          ),
          true
        )
        WHERE doc_id = %s
          AND raw -> 'raw_mineru_repair' ->> 'repair_id' = %s
        """
        with self._connect() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    sql,
                    (json.dumps(payload, ensure_ascii=False), doc_id, repair_id),
                )
            conn.commit()

    def _should_append_event(
        self,
        *,
        previous: dict[str, Any] | None,
        next_status: TaskStatus,
        next_progress: int,
        next_step: str,
    ) -> bool:
        """降低高频进度事件写入，同时保留关键状态流转。"""
        if previous is None:
            return True
        if next_status in {TaskStatus.COMPLETED, TaskStatus.FAILED}:
            return True
        previous_status = str(previous.get("status") or "")
        if previous_status != next_status.value:
            return True
        previous_progress = int(previous.get("latest_event_progress") or previous.get("progress") or 0)
        if abs(int(next_progress) - previous_progress) >= self._event_progress_step:
            return True
        latest_event_at = previous.get("latest_event_at")
        if latest_event_at and self._event_min_interval_seconds > 0:
            elapsed = (datetime.now(latest_event_at.tzinfo or timezone.utc) - latest_event_at).total_seconds()
            if elapsed >= self._event_min_interval_seconds:
                return True
        previous_step = str(previous.get("current_step") or "")
        if previous_step != next_step and self._event_min_interval_seconds <= 0:
            return True
        return False

    async def ensure_initialized(self) -> None:
        """确保表结构初始化只执行一次。"""
        async with self._init_lock:
            if self._initialized:
                return
            await asyncio.to_thread(self._init_schema_sync)
            self._initialized = True

    def _init_schema_sync(self) -> None:
        """同步初始化数据库表结构与索引。"""
        with self._connect() as conn:
            with conn.cursor() as cur:
                cur.execute(CREATE_TASKS_SQL)
                cur.execute(CREATE_EVENTS_SQL)
            conn.commit()

    async def upsert(self, task: TaskStatusData) -> None:
        """创建或更新任务记录。"""
        await self.ensure_initialized()
        await asyncio.to_thread(self._upsert_sync, task)

    def _upsert_sync(self, task: TaskStatusData) -> None:
        """同步写入任务主表并追加事件记录。"""
        payload = task.model_dump()
        with self._connect() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO ingestion_tasks (
                        task_id, doc_id, user_id, kb_id, status, progress, current_step, total_pages,
                        processed_pages, tree_node_count, error_message, error_code, retryable,
                        started_at, updated_at, completed_at, retry_count, eta_seconds, persisted_doc_id, persisted_doc_ids,
                        total_documents, processed_documents, callback_url, idempotency_key, is_temp, session_id
                    ) VALUES (
                        %(task_id)s, %(doc_id)s, %(user_id)s, %(kb_id)s, %(status)s, %(progress)s, %(current_step)s, %(total_pages)s,
                        %(processed_pages)s, %(tree_node_count)s, %(error_message)s, %(error_code)s, %(retryable)s,
                        %(started_at)s, %(updated_at)s, %(completed_at)s, %(retry_count)s, %(eta_seconds)s, %(persisted_doc_id)s,
                        %(persisted_doc_ids)s::jsonb, %(total_documents)s, %(processed_documents)s, %(callback_url)s, %(idempotency_key)s, %(is_temp)s, %(session_id)s
                    )
                    ON CONFLICT (task_id) DO UPDATE SET
                        doc_id = EXCLUDED.doc_id,
                        user_id = EXCLUDED.user_id,
                        kb_id = EXCLUDED.kb_id,
                        status = EXCLUDED.status,
                        progress = EXCLUDED.progress,
                        current_step = EXCLUDED.current_step,
                        total_pages = EXCLUDED.total_pages,
                        processed_pages = EXCLUDED.processed_pages,
                        tree_node_count = EXCLUDED.tree_node_count,
                        error_message = EXCLUDED.error_message,
                        error_code = EXCLUDED.error_code,
                        retryable = EXCLUDED.retryable,
                        started_at = EXCLUDED.started_at,
                        updated_at = EXCLUDED.updated_at,
                        completed_at = EXCLUDED.completed_at,
                        retry_count = EXCLUDED.retry_count,
                        eta_seconds = EXCLUDED.eta_seconds,
                        persisted_doc_id = EXCLUDED.persisted_doc_id,
                        persisted_doc_ids = EXCLUDED.persisted_doc_ids,
                        total_documents = EXCLUDED.total_documents,
                        processed_documents = EXCLUDED.processed_documents,
                        callback_url = EXCLUDED.callback_url,
                        idempotency_key = EXCLUDED.idempotency_key,
                        is_temp = EXCLUDED.is_temp,
                        session_id = EXCLUDED.session_id
                    """,
                    {**payload, "persisted_doc_ids": json.dumps(payload.get("persisted_doc_ids", []), ensure_ascii=False)},
                )
                cur.execute(
                    """
                    INSERT INTO ingestion_task_events(task_id, status, event_payload)
                    VALUES (%s, %s, %s::jsonb)
                    """,
                    (task.task_id, task.status.value, json.dumps(payload, ensure_ascii=False, default=str)),
                )
            conn.commit()

    async def get_by_task_id(self, task_id: str) -> TaskStatusData | None:
        """按任务 ID 查询任务状态。"""
        await self.ensure_initialized()
        row = await asyncio.to_thread(self._fetchone_sync, "task_id = %s", (task_id,))
        return self._row_to_model(row)

    async def get_latest_by_doc_id(self, doc_id: str) -> TaskStatusData | None:
        """按文档 ID 查询最新任务状态。"""
        await self.ensure_initialized()
        row = await asyncio.to_thread(
            self._fetchone_sync,
            "doc_id = %s ORDER BY updated_at DESC, created_at DESC LIMIT 1",
            (doc_id,),
            custom_where=True,
        )
        return self._row_to_model(row)

    async def get_latest_by_persisted_doc_id(
        self,
        persisted_doc_id: str,
        user_id: str,
        kb_id: str,
    ) -> TaskStatusData | None:
        """按用户、知识库与文档 UUID 查询最新任务状态。"""
        await self.ensure_initialized()
        row = await asyncio.to_thread(
            self._fetchone_sync,
            (
                "user_id = %s AND kb_id = %s AND "
                "(doc_id = %s OR persisted_doc_id = %s OR persisted_doc_ids @> %s::jsonb) "
                "ORDER BY updated_at DESC, created_at DESC LIMIT 1"
            ),
            (
                user_id,
                kb_id,
                persisted_doc_id,
                persisted_doc_id,
                json.dumps([persisted_doc_id], ensure_ascii=False),
            ),
            custom_where=True,
        )
        return self._row_to_model(row)

    async def get_by_biz_key(self, doc_id: str, user_id: str, kb_id: str) -> TaskStatusData | None:
        """按文档 ID、用户 ID 与知识库 ID 查询任务状态。"""
        await self.ensure_initialized()
        row = await asyncio.to_thread(
            self._fetchone_sync,
            "doc_id = %s AND user_id = %s AND kb_id = %s ORDER BY updated_at DESC, created_at DESC LIMIT 1",
            (doc_id, user_id, kb_id),
            custom_where=True,
        )
        return self._row_to_model(row)

    async def get_all_by_doc_id(
        self,
        doc_id: str,
        user_id: str | None = None,
        kb_id: str | None = None,
    ) -> list[TaskStatusData]:
        """查询文档关联的任务记录（可按用户与知识库过滤）。"""
        await self.ensure_initialized()
        rows = await asyncio.to_thread(self._fetchall_by_doc_sync, doc_id, user_id, kb_id)
        return [m for m in (self._row_to_model(row) for row in rows) if m is not None]

    async def get_by_idempotency_key(self, kb_id: str, user_id: str, idempotency_key: str) -> TaskStatusData | None:
        """按知识库、用户与幂等键查询历史任务。"""
        await self.ensure_initialized()
        row = await asyncio.to_thread(
            self._fetchone_sync,
            "kb_id = %s AND user_id = %s AND idempotency_key = %s ORDER BY updated_at DESC, created_at DESC LIMIT 1",
            (kb_id, user_id, idempotency_key),
            custom_where=True,
        )
        return self._row_to_model(row)

    def _fetchall_by_doc_sync(
        self,
        doc_id: str,
        user_id: str | None = None,
        kb_id: str | None = None,
    ) -> list[dict[str, Any]]:
        """同步查询指定文档的任务记录（可按用户与知识库过滤）。"""
        conditions = ["doc_id = %s"]
        params: list[Any] = [doc_id]
        if user_id:
            conditions.append("user_id = %s")
            params.append(user_id)
        if kb_id:
            conditions.append("kb_id = %s")
            params.append(kb_id)
        query = f"""
            SELECT task_id, doc_id, user_id, kb_id, status, progress, current_step, total_pages, processed_pages,
                   tree_node_count, error_message, error_code, retryable, started_at, updated_at,
                   completed_at, retry_count, eta_seconds, persisted_doc_id, persisted_doc_ids,
                   total_documents, processed_documents, callback_url, idempotency_key, is_temp, session_id
            FROM ingestion_tasks
            WHERE {' AND '.join(conditions)}
            ORDER BY updated_at DESC, created_at DESC
        """
        with self._connect() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute(query, tuple(params))
                rows = cur.fetchall() or []
                return [dict(r) for r in rows]

    def _fetchone_sync(
        self,
        where_expr: str,
        params: tuple[Any, ...],
        *,
        custom_where: bool = False,
    ) -> dict[str, Any] | None:
        """按条件同步查询单条任务记录。"""
        if custom_where:
            query = f"""
                SELECT task_id, doc_id, user_id, kb_id, status, progress, current_step, total_pages, processed_pages,
                       tree_node_count, error_message, error_code, retryable, started_at, updated_at,
                       completed_at, retry_count, eta_seconds, persisted_doc_id, persisted_doc_ids,
                       total_documents, processed_documents, callback_url, idempotency_key, is_temp, session_id
                FROM ingestion_tasks
                WHERE {where_expr}
            """
        else:
            query = f"""
                SELECT task_id, doc_id, user_id, kb_id, status, progress, current_step, total_pages, processed_pages,
                       tree_node_count, error_message, error_code, retryable, started_at, updated_at,
                       completed_at, retry_count, eta_seconds, persisted_doc_id, persisted_doc_ids,
                       total_documents, processed_documents, callback_url, idempotency_key, is_temp, session_id
                FROM ingestion_tasks
                WHERE {where_expr}
                LIMIT 1
            """
        with self._connect() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute(query, params)
                row = cur.fetchone()
                return dict(row) if row else None

    async def update_status(
        self,
        task_id: str,
        *,
        status: TaskStatus,
        progress: int,
        current_step: str,
        total_pages: int | None = None,
        processed_pages: int | None = None,
        tree_node_count: int | None = None,
        error_message: str | None = None,
        error_code: str | None = None,
        retryable: bool | None = None,
        eta_seconds: int | None = None,
        persisted_doc_id: str | None = None,
        persisted_doc_ids: list[str] | None = None,
        total_documents: int | None = None,
        processed_documents: int | None = None,
    ) -> TaskStatusData | None:
        """更新数据库中的任务状态并写入事件流。"""
        await self.ensure_initialized()
        updated = await asyncio.to_thread(
            self._update_status_sync,
            task_id,
            status=status,
            progress=progress,
            current_step=current_step,
            total_pages=total_pages,
            processed_pages=processed_pages,
            tree_node_count=tree_node_count,
            error_message=error_message,
            error_code=error_code,
            retryable=retryable,
            eta_seconds=eta_seconds,
            persisted_doc_id=persisted_doc_id,
            persisted_doc_ids=persisted_doc_ids,
            total_documents=total_documents,
            processed_documents=processed_documents,
        )
        return updated

    def _update_status_sync(
        self,
        task_id: str,
        *,
        status: TaskStatus,
        progress: int,
        current_step: str,
        total_pages: int | None = None,
        processed_pages: int | None = None,
        tree_node_count: int | None = None,
        error_message: str | None = None,
        error_code: str | None = None,
        retryable: bool | None = None,
        eta_seconds: int | None = None,
        persisted_doc_id: str | None = None,
        persisted_doc_ids: list[str] | None = None,
        total_documents: int | None = None,
        processed_documents: int | None = None,
    ) -> TaskStatusData | None:
        """同步更新任务状态并写入状态事件。"""
        with self._connect() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute(
                    """
                    SELECT
                        status,
                        progress,
                        current_step,
                        (
                            SELECT created_at
                            FROM ingestion_task_events
                            WHERE task_id = %s
                            ORDER BY created_at DESC
                            LIMIT 1
                        ) AS latest_event_at,
                        (
                            SELECT event_payload->>'progress'
                            FROM ingestion_task_events
                            WHERE task_id = %s
                            ORDER BY created_at DESC
                            LIMIT 1
                        ) AS latest_event_progress
                    FROM ingestion_tasks
                    WHERE task_id = %s
                    FOR UPDATE
                    """,
                    (task_id, task_id, task_id),
                )
                previous_row = cur.fetchone()
                previous = dict(previous_row) if previous_row else None
                next_progress = max(0, min(100, int(progress)))
                append_event = self._should_append_event(
                    previous=previous,
                    next_status=status,
                    next_progress=next_progress,
                    next_step=current_step,
                )
                cur.execute(
                    """
                    UPDATE ingestion_tasks
                    SET
                        status = %(status)s,
                        progress = %(progress)s,
                        current_step = %(current_step)s,
                        total_pages = COALESCE(%(total_pages)s, total_pages),
                        processed_pages = COALESCE(%(processed_pages)s, processed_pages),
                        tree_node_count = COALESCE(%(tree_node_count)s, tree_node_count),
                        error_message = %(error_message)s,
                        error_code = %(error_code)s,
                        retryable = COALESCE(%(retryable)s, retryable),
                        eta_seconds = COALESCE(%(eta_seconds)s, eta_seconds),
                        persisted_doc_id = COALESCE(%(persisted_doc_id)s, persisted_doc_id),
                        persisted_doc_ids = COALESCE(%(persisted_doc_ids)s::jsonb, persisted_doc_ids),
                        total_documents = COALESCE(%(total_documents)s, total_documents),
                        processed_documents = COALESCE(%(processed_documents)s, processed_documents),
                        updated_at = NOW(),
                        completed_at = CASE
                            WHEN %(status)s IN ('completed', 'failed') AND completed_at IS NULL THEN NOW()
                            ELSE completed_at
                        END
                    WHERE task_id = %(task_id)s
                    RETURNING task_id, doc_id, user_id, kb_id, status, progress, current_step, total_pages, processed_pages,
                              tree_node_count, error_message, error_code, retryable, started_at, updated_at,
                              completed_at, retry_count, eta_seconds, persisted_doc_id, persisted_doc_ids,
                              total_documents, processed_documents, callback_url, idempotency_key, is_temp, session_id
                    """,
                    {
                        "task_id": task_id,
                        "status": status.value,
                        "progress": next_progress,
                        "current_step": current_step,
                        "total_pages": total_pages,
                        "processed_pages": processed_pages,
                        "tree_node_count": tree_node_count,
                        "error_message": error_message,
                        "error_code": error_code,
                        "retryable": retryable,
                        "eta_seconds": eta_seconds,
                        "persisted_doc_id": persisted_doc_id,
                        "persisted_doc_ids": None if persisted_doc_ids is None else json.dumps(persisted_doc_ids, ensure_ascii=False),
                        "total_documents": total_documents,
                        "processed_documents": processed_documents,
                    },
                )
                row = cur.fetchone()
                if row is None:
                    conn.commit()
                    return None
                payload = dict(row)
                if append_event:
                    cur.execute(
                        """
                        INSERT INTO ingestion_task_events(task_id, status, event_payload)
                        VALUES (%s, %s, %s::jsonb)
                        """,
                        (task_id, status.value, json.dumps(payload, ensure_ascii=False, default=str)),
                    )
            conn.commit()
        return self._row_to_model(payload)

    async def remove_by_doc_id(
        self,
        doc_id: str,
        user_id: str | None = None,
        kb_id: str | None = None,
    ) -> int:
        """删除文档关联任务记录（可按用户与知识库过滤）。"""
        await self.ensure_initialized()
        return await asyncio.to_thread(self._remove_by_doc_sync, doc_id, user_id, kb_id)

    async def get_temp_by_session(self, session_id: str, kb_id: str | None = None) -> list[TaskStatusData]:
        """查询某个会话下的临时文档任务。"""
        await self.ensure_initialized()
        return await asyncio.to_thread(self._get_temp_by_session_sync, session_id, kb_id)

    def _get_temp_by_session_sync(self, session_id: str, kb_id: str | None = None) -> list[TaskStatusData]:
        """同步查询某个会话下的临时文档任务。"""
        base_sql = """
            SELECT task_id, doc_id, user_id, kb_id, status, progress, current_step, total_pages, processed_pages,
                   tree_node_count, error_message, error_code, retryable, started_at, updated_at,
                   completed_at, retry_count, eta_seconds, persisted_doc_id, persisted_doc_ids,
                   total_documents, processed_documents, callback_url, idempotency_key, is_temp, session_id
            FROM ingestion_tasks
            WHERE is_temp = TRUE AND session_id = %s
        """
        params: list[Any] = [session_id]
        if kb_id:
            base_sql += " AND kb_id = %s"
            params.append(kb_id)
        base_sql += " ORDER BY updated_at DESC, created_at DESC"
        out: list[TaskStatusData] = []
        with self._connect() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute(base_sql, tuple(params))
                for row in cur.fetchall():
                    model = self._row_to_model(dict(row))
                    if model is not None:
                        out.append(model)
        return out

    def _remove_by_doc_sync(self, doc_id: str, user_id: str | None = None, kb_id: str | None = None) -> int:
        """同步删除文档关联任务与事件数据（可按用户与知识库过滤）。"""
        conditions = ["doc_id = %s"]
        params: list[Any] = [doc_id]
        if user_id:
            conditions.append("user_id = %s")
            params.append(user_id)
        if kb_id:
            conditions.append("kb_id = %s")
            params.append(kb_id)
        where_sql = " AND ".join(conditions)
        with self._connect() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    f"DELETE FROM ingestion_task_events WHERE task_id IN (SELECT task_id FROM ingestion_tasks WHERE {where_sql})",
                    tuple(params),
                )
                cur.execute(f"DELETE FROM ingestion_tasks WHERE {where_sql}", tuple(params))
                deleted = cur.rowcount or 0
            conn.commit()
        return int(deleted)

    @staticmethod
    def _row_to_model(row: dict[str, Any] | None) -> TaskStatusData | None:
        """把数据库行转换为 `TaskStatusData` 模型。"""
        if row is None:
            return None
        normalized = dict(row)
        if normalized.get("persisted_doc_ids") is None:
            normalized["persisted_doc_ids"] = []
        for key in ("started_at", "updated_at", "completed_at"):
            value = normalized.get(key)
            if value is None:
                continue
            try:
                if hasattr(value, "astimezone"):
                    value = value.astimezone(timezone.utc)
                normalized[key] = value.isoformat().replace("+00:00", "Z")
            except Exception:
                normalized[key] = str(value)
        raw_status = normalized.get("status")
        if isinstance(raw_status, TaskStatus):
            normalized["status"] = raw_status
        else:
            status_text = str(raw_status or "").strip()
            # 向后兼容：旧数据行里可能存的是 "TaskStatus.PARSING"。
            if status_text.startswith("TaskStatus."):
                status_text = status_text.split(".", 1)[1].lower()
            normalized["status"] = TaskStatus(status_text)
        return TaskStatusData.model_validate(normalized)

class ProgressReporter(Protocol):
    """执行器上报进度的回调协议。"""
    # 该回调由执行器在各阶段调用，用于上报前端可见的进度与状态。
    async def __call__(
        self,
        *,
        status: str,
        progress: int,
        current_step: str,
        total_pages: int | None = None,
        processed_pages: int | None = None,
        tree_node_count: int | None = None,
        eta_seconds: int | None = None,
        error_message: str | None = None,
        error_code: str | None = None,
        retryable: bool | None = None,
        persisted_doc_id: str | None = None,
        persisted_doc_ids: list[str] | None = None,
        total_documents: int | None = None,
        processed_documents: int | None = None,
    ) -> None: ...


class IngestionAlgorithmRunner(Protocol):
    """入库算法执行器协议。"""
    async def run(self, request: IngestRequest, report: ProgressReporter) -> None: ...


class RawMineruRepairTerminalError(RuntimeError):
    """MinerU repair failed in a way that should not be retried."""


class DocumentAssistantRunner:
    """入库执行器。

    使用独立多格式解析器生成结构化文档，并写入 PostgreSQL。
    """
    # 进度分段约定：
    # 0~44: 预热与解析阶段（心跳推进）
    # 45~88: 构建树结构阶段（心跳推进）
    # 90: 持久化阶段
    # 100: 完成
    PARSING_PROGRESS_START = 5
    EXTRACTING_PROGRESS_START = 25
    STRUCTURING_PROGRESS_START = 45
    STRUCTURING_PROGRESS_MAX = 88
    STORING_PROGRESS = 90
    COMPLETED_PROGRESS = 100

    def __init__(
        self,
        *,
        workspace: str | None = None,
        model: str | None = None,
        retrieve_model: str | None = None,
        adapter: DocumentAssistantAdapter | None = None,
        mineru_parser: Callable[..., Awaitable[dict[str, Any]]] | None = None,
    ) -> None:
        """初始化对象并准备运行所需的配置与状态。"""
        self._logger = logging.getLogger(__name__)
        self._adapter = adapter or DocumentAssistantAdapter()
        self._mineru_parser = mineru_parser or parse_document_with_mineru_async
        env_workspace = WORKSPACE_DIR
        self.workspace = workspace or env_workspace or str((Path(__file__).resolve().parent.parent / "workspace").resolve())
        self._allowed_input_schemes = set(ALLOWED_INPUT_SCHEMES)
        self._allow_local_input = "local" in self._allowed_input_schemes
        resolved_model = (
            model
            or str(os.getenv("ARK_MODEL", "") or "").strip()
            or None
        )
        self.model = resolved_model
        self.retrieve_model = (
            retrieve_model
            or str(os.getenv("ARK_RETRIEVE_MODEL", "") or "").strip()
            or self.model
        )

    async def _repair_raw_mineru(
        self,
        *,
        client: Any,
        file_path: Path,
        doc_id: str,
        repair_id: str | None = None,
    ) -> None:
        """Reparse with MinerU and patch only the stored raw_mineru field."""
        config = RuntimeConfig()
        if str(config.parser_backend).strip().lower() != "mineru":
            raise RuntimeError("raw_mineru repair requires the MinerU parser backend")
        payload = await self._mineru_parser(
            file_path,
            api_url=config.mineru_api_url,
            timeout_seconds=config.mineru_timeout_seconds,
            retry_times=config.mineru_retry_times,
            retry_backoff_base=config.mineru_retry_backoff_base,
            poll_interval_seconds=config.mineru_poll_interval_seconds,
            backend=config.mineru_backend,
            parse_method=config.mineru_parse_method,
            lang=config.mineru_lang,
            use_async_tasks=config.mineru_use_async_tasks,
            client_concurrency=config.mineru_client_concurrency,
            weak_heading_split_enabled=config.weak_heading_split_enabled,
            weak_heading_split_min_chars=config.weak_heading_split_min_chars,
            weak_heading_split_max_chars=config.weak_heading_split_max_chars,
        )
        raw_mineru = payload.get("raw_mineru") if isinstance(payload, dict) else None
        if not is_complete_raw_mineru(raw_mineru):
            raise RawMineruRepairTerminalError("MinerU repair returned an incomplete raw payload")
        store = getattr(client, "_store", None)
        if store is None or not hasattr(store, "patch_raw_mineru"):
            raise RuntimeError("storage backend does not support raw_mineru repair")
        await asyncio.to_thread(
            store.patch_raw_mineru,
            doc_id,
            raw_mineru,
            repair_id=repair_id,
        )

    async def run(self, request: IngestRequest, report: ProgressReporter) -> None:
        """执行单文档/批量文档入库流程并持续上报进度。"""
        inputs = request.iter_documents()
        total_docs = max(1, len(inputs))
        task_temp_dir = Path(tempfile.mkdtemp(prefix=f"ingest_{request.doc_id}_{uuid.uuid4().hex[:8]}_")).resolve()
        self._logger.info(
            "ingestion_task_start doc_id=%s kb_id=%s is_temp=%s session_id=%s total_docs=%s temp_dir=%s",
            request.doc_id,
            request.kb_id,
            bool(request.is_temp),
            request.session_id,
            total_docs,
            task_temp_dir,
        )
        await report(
            status="parsing",
            progress=1,
            current_step="准备输入文件",
            total_documents=total_docs,
            processed_documents=0,
            eta_seconds=max(30, 20 * total_docs),
        )
        warmup_done = asyncio.Event()
        warmup_heartbeat = asyncio.create_task(
            # 预热心跳：避免 queued 后长时间停在 0。
            self._heartbeat_warmup_progress(
                report=report,
                stop_event=warmup_done,
                total_documents=total_docs,
                processed_documents=0,
                start_progress=1,
                max_progress=self.STRUCTURING_PROGRESS_START - 1,
                tick_seconds=1.0,
            )
        )
        local_files: list[Path] = []
        source_metadata: list[dict[str, str]] = []
        try:
            local_files, source_metadata = await self._resolve_input_files(inputs, temp_dir=task_temp_dir)
            client = self._build_client(user_id=request.user_id, kb_id=request.kb_id, session_id=request.session_id)
            if request.raw_mineru_repair_doc_ids:
                if len(local_files) != 1 or request.raw_mineru_repair_doc_ids != {request.doc_id}:
                    raise RuntimeError("raw_mineru repair only supports one existing document")
                await self._stop_heartbeat_task(warmup_done, warmup_heartbeat)
                await report(
                    status="extracting",
                    progress=self.EXTRACTING_PROGRESS_START,
                    current_step="repairing MinerU raw output",
                    total_documents=1,
                    processed_documents=0,
                    eta_seconds=30,
                )
                await report(
                    status="structuring",
                    progress=self.STRUCTURING_PROGRESS_START,
                    current_step="validating MinerU raw output",
                    total_documents=1,
                    processed_documents=0,
                    eta_seconds=25,
                )
                repair_store = getattr(client, "_store", None)
                if (
                    request.raw_mineru_repair_id
                    and repair_store is not None
                    and hasattr(repair_store, "update_raw_mineru_repair")
                ):
                    await asyncio.to_thread(
                        repair_store.update_raw_mineru_repair,
                        doc_id=request.doc_id,
                        repair_id=request.raw_mineru_repair_id,
                        updates={
                            "status": "processing",
                            "updated_at": datetime.now(timezone.utc).isoformat(),
                            "lease_expires_at": (
                                datetime.now(timezone.utc) + timedelta(seconds=1200)
                            ).isoformat(),
                        },
                    )
                await self._repair_raw_mineru(
                    client=client,
                    file_path=local_files[0],
                    doc_id=request.doc_id,
                    repair_id=request.raw_mineru_repair_id,
                )
                await report(
                    status="storing",
                    progress=self.STORING_PROGRESS,
                    current_step="patching MinerU raw output",
                    total_documents=1,
                    processed_documents=1,
                    eta_seconds=1,
                )
                await report(
                    status="completed",
                    progress=self.COMPLETED_PROGRESS,
                    current_step=f"MinerU raw repair completed ({request.doc_id})",
                    total_documents=1,
                    processed_documents=1,
                    eta_seconds=0,
                    persisted_doc_id=request.doc_id,
                    persisted_doc_ids=[request.doc_id],
                )
                self._logger.info(
                    "ingestion_raw_mineru_repair_completed doc_id=%s kb_id=%s",
                    request.doc_id,
                    request.kb_id,
                )
                return
            runtime = None
            options = self._adapter.build_indexing_options(
                resolved_overlap_ratio=request.config.resolved_overlap_ratio(),
                summary_enabled=bool(request.config.summary_enabled),
                table_parse_mode=request.config.resolved_table_parse_mode(),
                node_max_tokens=int(request.config.node_max_tokens),
                max_tree_depth=int(request.config.max_tree_depth),
                enable_ocr=bool(request.config.enable_ocr),
            )
            if len(local_files) <= 1:
                local_file = local_files[0]
                # 单文档模式下尽量提前拿到页数，用于中间进度展示。
                total_pages = await asyncio.to_thread(self._safe_get_document_pages, local_file)
                estimated_total_nodes = self._estimate_total_nodes(total_pages=total_pages)
                await self._stop_heartbeat_task(warmup_done, warmup_heartbeat)
                await report(
                    status="extracting",
                    progress=self.EXTRACTING_PROGRESS_START,
                    current_step="extracting content",
                    total_pages=total_pages,
                    processed_pages=0,
                    tree_node_count=0,
                    total_documents=1,
                    processed_documents=0,
                    eta_seconds=35 if total_pages else 20,
                )
                await report(
                    status="structuring",
                    progress=self.STRUCTURING_PROGRESS_START,
                    current_step="building pageindex tree",
                    total_pages=total_pages,
                    processed_pages=0,
                    tree_node_count=0,
                    total_documents=1,
                    processed_documents=0,
                    eta_seconds=25 if total_pages else 15,
                )
                structuring_done = asyncio.Event()
                structuring_heartbeat = asyncio.create_task(
                    # structuring 心跳：持续刷新 progress/processed_pages/tree_node_count。
                    self._heartbeat_structuring_progress(
                        report=report,
                        stop_event=structuring_done,
                        current_step="building pageindex tree",
                        total_pages=total_pages,
                        estimated_total_nodes=estimated_total_nodes,
                        total_documents=1,
                        processed_documents=0,
                        start_progress=self.STRUCTURING_PROGRESS_START,
                        max_progress=self.STRUCTURING_PROGRESS_MAX,
                        tick_seconds=3.0,
                    )
                )
                try:
                    payload = await self._run_real_pipeline_or_fallback(
                        client=client,
                        file_path=local_file,
                        options=options,
                        runtime=runtime,
                        source_metadata=(source_metadata[0] if source_metadata else None),
                    )
                finally:
                    structuring_done.set()
                    structuring_heartbeat.cancel()
                    with contextlib.suppress(asyncio.CancelledError):
                        await structuring_heartbeat
                if request.is_temp:
                    payload["is_temporary"] = True
                    payload["session_id"] = request.session_id
                payload["id"] = request.doc_id
                payload["doc_id"] = request.doc_id

                pages = int(payload.get("page_count") or total_pages or 0)
                nodes = len(payload.get("structure") or [])
                await report(
                    status="storing",
                    progress=self.STORING_PROGRESS,
                    current_step="storing into backend",
                    total_pages=pages or None,
                    processed_pages=pages,
                    tree_node_count=nodes,
                    total_documents=1,
                    processed_documents=1,
                    eta_seconds=8,
                )
                persisted_doc_id = await asyncio.to_thread(self._persist_payload_or_doc_id, client, payload)
                await report(
                    status="completed",
                    progress=self.COMPLETED_PROGRESS,
                    current_step=f"ingestion completed ({persisted_doc_id})",
                    total_pages=pages or None,
                    processed_pages=pages,
                    tree_node_count=nodes,
                    total_documents=1,
                    processed_documents=1,
                    eta_seconds=0,
                    persisted_doc_id=str(persisted_doc_id),
                    persisted_doc_ids=[str(persisted_doc_id)],
                )
                self._logger.info(
                    "ingestion_task_completed doc_id=%s kb_id=%s persisted_doc_id=%s pages=%s nodes=%s",
                    request.doc_id,
                    request.kb_id,
                    persisted_doc_id,
                    pages,
                    nodes,
                )
                return

            await self._stop_heartbeat_task(warmup_done, warmup_heartbeat)
            await report(
                status="extracting",
                progress=self.EXTRACTING_PROGRESS_START,
                current_step=f"resolving content from {len(local_files)} documents",
                total_documents=len(local_files),
                processed_documents=0,
                eta_seconds=max(35, 20 * len(local_files)),
            )
            await report(
                status="structuring",
                progress=self.STRUCTURING_PROGRESS_START,
                current_step=f"批量索引 {len(local_files)} 个文档",
                total_documents=len(local_files),
                processed_documents=0,
                eta_seconds=max(20, 15 * len(local_files)),
            )
            persisted_doc_ids = await self._run_batch_pipeline(
                client=client,
                file_paths=local_files,
                target_doc_ids=request.doc_ids,
                source_metadata=source_metadata,
                options=options,
                runtime=runtime,
                report=report,
            )
            await report(
                status="structuring",
                progress=self.STRUCTURING_PROGRESS_MAX,
                current_step=f"indexed {len(persisted_doc_ids)} documents",
                total_documents=len(local_files),
                processed_documents=len(persisted_doc_ids),
                eta_seconds=10,
            )
            await report(
                status="storing",
                progress=self.STORING_PROGRESS,
                current_step=f"persisted {len(persisted_doc_ids)} documents",
                total_documents=len(local_files),
                processed_documents=len(persisted_doc_ids),
                eta_seconds=8,
            )
            await report(
                status="completed",
                progress=self.COMPLETED_PROGRESS,
                current_step=f"批量入库完成（{len(persisted_doc_ids)} 个文档）",
                total_documents=len(local_files),
                processed_documents=len(persisted_doc_ids),
                eta_seconds=0,
                persisted_doc_id=str(persisted_doc_ids[0]),
                persisted_doc_ids=[str(x) for x in persisted_doc_ids],
            )
            self._logger.info(
                "ingestion_task_completed doc_id=%s kb_id=%s persisted_doc_ids=%s total_docs=%s",
                request.doc_id,
                request.kb_id,
                [str(x) for x in persisted_doc_ids],
                len(local_files),
            )
        except requests.RequestException as exc:
            self._logger.exception("ingestion_task_download_failed doc_id=%s kb_id=%s", request.doc_id, request.kb_id)
            await report(status="failed", progress=99, current_step="文件下载失败", error_message=str(exc), error_code="FILE_DOWNLOAD_FAILED", retryable=True, eta_seconds=0)
            raise
        except DuplicateDocumentError as exc:
            self._logger.info("ingestion_task_duplicate_document doc_id=%s kb_id=%s", request.doc_id, request.kb_id)
            await report(status="failed", progress=99, current_step="重复文档", error_message=str(exc), error_code="DUPLICATE_DOCUMENT", retryable=False, eta_seconds=0)
            raise
        except Exception as exc:
            self._logger.exception("ingestion_task_failed doc_id=%s kb_id=%s", request.doc_id, request.kb_id)
            await report(status="failed", progress=99, current_step="入库失败", error_message=str(exc), error_code="INGESTION_EXECUTION_FAILED", retryable=True, eta_seconds=0)
            raise
        finally:
            await self._stop_heartbeat_task(warmup_done, warmup_heartbeat)
            with contextlib.suppress(Exception):
                if task_temp_dir.exists():
                    shutil.rmtree(task_temp_dir, ignore_errors=True)
                    self._logger.info(
                        "ingestion_task_temp_dir_cleaned doc_id=%s kb_id=%s temp_dir=%s",
                        request.doc_id,
                        request.kb_id,
                        task_temp_dir,
                    )
            submit_temp_dir = Path(request._submit_temp_dir).resolve() if request._submit_temp_dir else None
            with contextlib.suppress(Exception):
                if submit_temp_dir is not None and submit_temp_dir.exists():
                    shutil.rmtree(submit_temp_dir, ignore_errors=True)
                    self._logger.info(
                        "ingestion_submit_temp_dir_cleaned doc_id=%s kb_id=%s temp_dir=%s",
                        request.doc_id,
                        request.kb_id,
                        submit_temp_dir,
                    )

    @staticmethod
    async def _heartbeat_structuring_progress(
        *,
        report: ProgressReporter,
        stop_event: asyncio.Event,
        current_step: str,
        total_pages: int | None,
        estimated_total_nodes: int | None,
        total_documents: int,
        processed_documents: int,
        start_progress: int = 45,
        max_progress: int = 88,
        tick_seconds: float = 3.0,
    ) -> None:
        """在 structuring 阶段定期上报心跳进度，避免进度长时间停滞。

        同时按进度比例推算 processed_pages/tree_node_count，提升前端可观测性。
        """
        progress = int(start_progress)
        while not stop_event.is_set():
            try:
                await asyncio.wait_for(stop_event.wait(), timeout=tick_seconds)
                break
            except asyncio.TimeoutError:
                progress = min(int(max_progress), progress + 1)
                processed_pages_now, tree_nodes_now = DocumentAssistantRunner._derive_structuring_metrics(
                    progress=progress,
                    start_progress=int(start_progress),
                    max_progress=int(max_progress),
                    total_pages=total_pages,
                    estimated_total_nodes=estimated_total_nodes,
                )
                await report(
                    status="structuring",
                    progress=progress,
                    current_step=current_step,
                    total_pages=total_pages,
                    processed_pages=processed_pages_now,
                    tree_node_count=tree_nodes_now,
                    total_documents=total_documents,
                    processed_documents=processed_documents,
                    eta_seconds=max(5, (int(max_progress) - progress) * 2),
                )

    @staticmethod
    async def _heartbeat_warmup_progress(
        *,
        report: ProgressReporter,
        stop_event: asyncio.Event,
        total_documents: int,
        processed_documents: int,
        start_progress: int = 1,
        max_progress: int = 44,
        tick_seconds: float = 1.0,
    ) -> None:
        """在 0~45 前持续更新进度，避免 queued 后长时间无变化。"""
        progress = max(1, int(start_progress))
        upper = max(progress, int(max_progress))
        while not stop_event.is_set():
            try:
                await asyncio.wait_for(stop_event.wait(), timeout=tick_seconds)
                break
            except asyncio.TimeoutError:
                progress = min(upper, progress + 1)
                # 预热区间内按阶段映射到 parsing/extracting，保证状态迁移合法。
                if progress < 15:
                    status = "parsing"
                    step = "准备输入文件"
                elif progress < 30:
                    status = "parsing"
                    step = "parsing document"
                else:
                    status = "extracting"
                    step = "extracting content"
                await report(
                    status=status,
                    progress=progress,
                    current_step=step,
                    total_documents=total_documents,
                    processed_documents=processed_documents,
                    eta_seconds=max(8, (upper - progress) * 2),
                )

    @staticmethod
    async def _stop_heartbeat_task(stop_event: asyncio.Event, task: asyncio.Task | None) -> None:
        """安全停止进度心跳任务。"""
        if task is None:
            return
        if not stop_event.is_set():
            stop_event.set()
        if task.done():
            with contextlib.suppress(asyncio.CancelledError):
                await task
            return
        task.cancel()
        with contextlib.suppress(asyncio.CancelledError):
            await task

    def _build_client(self, *, user_id: str, kb_id: str, session_id: str | None) -> Any:
        """按配置构建 PageIndex 客户端实例。"""
        from app.storage import PageIndexClient

        dsn = str(os.getenv("POSTGRES_DSN", "")).strip()
        if not dsn:
            raise RuntimeError("POSTGRES_DSN is required for ingestion client")
        return PageIndexClient(
            workspace=self.workspace,
            model=self.model,
            retrieve_model=self.retrieve_model,
            storage_backend="postgres",
            postgres_dsn=dsn,
            storage_fallback_to_file=False,
            user_id=str(user_id or os.getenv("DA_USER_ID", "") or os.getenv("PAGEINDEX_USER_ID", "") or "system"),
            kb_id=str(kb_id or os.getenv("DA_KB_ID", "") or os.getenv("PAGEINDEX_KB_ID", "") or "default"),
            session_id=session_id,
        )

    async def _run_real_pipeline_or_fallback(
        self,
        *,
        client: Any,
        file_path: Path,
        options: Any,
        runtime: Any,
        source_metadata: dict[str, Any] | None = None,
    ) -> dict:
        """优先走真实索引链路，失败时回退到直接索引。"""
        try:
            # 基于文件指纹做去重，避免同一文件重复索引。
            fingerprint = await self._adapter.compute_file_fingerprint(file_path)
            cached_doc_id = await self._adapter.find_cached_doc_id(client, file_path, fingerprint)
            if cached_doc_id:
                cached_doc = dict(client.documents.get(str(cached_doc_id), {}))
                needs_full = not isinstance(cached_doc.get("pages"), list) or not isinstance(cached_doc.get("structure"), list)
                if not cached_doc:
                    store = getattr(client, "_store", None)
                    if store is not None:
                        with contextlib.suppress(Exception):
                            full = store.load_full_doc(str(cached_doc_id))
                            if isinstance(full, dict):
                                cached_doc = dict(full)
                elif needs_full:
                    store = getattr(client, "_store", None)
                    if store is not None:
                        with contextlib.suppress(Exception):
                            full = store.load_full_doc(str(cached_doc_id))
                            if isinstance(full, dict):
                                cached_doc = dict(full)
                if cached_doc:
                    # 统一补齐 id 字段，保证后续接口返回结构稳定。
                    cached_doc.setdefault("id", str(cached_doc_id))
                    cached_doc.setdefault("doc_id", str(cached_doc_id))
                    self._logger.info(
                        "dedupe_hit_path_fingerprint file=%s doc_id=%s",
                        file_path.name,
                        cached_doc_id,
                    )
                    return self._adapter.apply_source_metadata(cached_doc, source_metadata)
            normalized = await self._adapter.load_or_index_document(
                model=client.model,
                retrieve_model=client.retrieve_model,
                file_path=file_path,
                options=options,
                runtime=runtime,
                fingerprint=fingerprint,
                source_metadata=source_metadata,
            )
            if normalized is None:
                raise RuntimeError("document indexing failed")
            return dict(normalized.payload)
        except (ImportError, ModuleNotFoundError) as exc:
            # Fallback only for missing optional deps; other runtime failures should surface.
            self._logger.warning("索引流水线不可用，回退到直接索引：%s", exc)
            doc_id = await asyncio.to_thread(client.index, str(file_path))
            doc = dict(client.documents.get(doc_id, {}))
            if not doc:
                raise RuntimeError("回退索引失败")
            doc["doc_id"] = doc_id
            return self._adapter.apply_source_metadata(doc, source_metadata)

    async def _run_batch_pipeline(
        self,
        *,
        client: Any,
        file_paths: list[Path],
        target_doc_ids: list[str] | None = None,
        options: Any,
        runtime: Any,
        source_metadata: list[dict[str, Any]] | None = None,
        report: ProgressReporter | None = None,
    ) -> list[str]:
        """执行批量文档入库流程并返回文档 ID 列表。"""
        try:
            index_task = asyncio.create_task(
                self._adapter.index_documents_async(
                    client,
                    file_paths,
                    target_doc_ids=target_doc_ids,
                    source_metadata=source_metadata,
                    options=options,
                    runtime=runtime,
                )
            )
            total = max(1, len(file_paths))
            last_completed = -1
            try:
                while not index_task.done():
                    await asyncio.sleep(1.0)
                    if report is None:
                        continue
                    stats = getattr(runtime, "stats", {}) if runtime is not None else {}
                    success = max(0, int(stats.get("index_success", 0) or 0))
                    failed = max(0, int(stats.get("index_fail", 0) or 0))
                    completed = min(total, success + failed)
                    if completed == last_completed:
                        continue
                    last_completed = completed
                    ratio = completed / total
                    progress_min = self.STRUCTURING_PROGRESS_START
                    progress_max = self.STRUCTURING_PROGRESS_MAX
                    progress_span = max(0, progress_max - progress_min)
                    await report(
                        status="structuring",
                        progress=min(progress_max, progress_min + int(ratio * progress_span)),
                        current_step=f"indexed {completed}/{total} documents",
                        total_documents=total,
                        processed_documents=completed,
                        eta_seconds=max(8, int((total - completed) * 12)),
                    )
                doc_ids, failed_docs = await index_task
            except asyncio.CancelledError:
                index_task.cancel()
                with contextlib.suppress(asyncio.CancelledError):
                    await index_task
                raise
            if failed_docs:
                raise RuntimeError(f"{len(failed_docs)} 个文档在批量入库时失败")
            if not doc_ids:
                raise RuntimeError("批量入库未产出任何文档")
            return [str(doc_id) for doc_id in doc_ids]
        except (ImportError, ModuleNotFoundError) as exc:
            self._logger.warning("批量索引流水线不可用，回退到顺序索引：%s", exc)
            doc_ids: list[str] = []
            target_doc_ids = [str(item) for item in (target_doc_ids or [])]
            for index, file_path in enumerate(file_paths):
                doc_id = await asyncio.to_thread(client.index, str(file_path))
                if index < len(target_doc_ids) and target_doc_ids[index]:
                    target_doc_id = target_doc_ids[index]
                    doc = dict(getattr(client, "documents", {}).get(str(doc_id), {}))
                    if doc:
                        doc["id"] = target_doc_id
                        doc["doc_id"] = target_doc_id
                        client.documents[target_doc_id] = doc
                        await asyncio.to_thread(client._save_doc, target_doc_id)
                        doc_id = target_doc_id
                doc_ids.append(str(doc_id))
                if report is not None:
                    ratio = len(doc_ids) / max(1, len(file_paths))
                    progress_min = self.STRUCTURING_PROGRESS_START
                    progress_max = self.STRUCTURING_PROGRESS_MAX
                    progress_span = max(0, progress_max - progress_min)
                    await report(
                        status="structuring",
                        progress=min(progress_max, progress_min + int(ratio * progress_span)),
                        current_step=f"indexed {len(doc_ids)}/{len(file_paths)} documents",
                        total_documents=len(file_paths),
                        processed_documents=len(doc_ids),
                        eta_seconds=max(8, int((len(file_paths) - len(doc_ids)) * 12)),
                    )
            if not doc_ids:
                raise RuntimeError("batch 回退索引失败")
            return doc_ids

    def _persist_payload_or_doc_id(self, client: Any, payload: dict) -> str:
        """将索引结果持久化并返回文档 ID。"""
        try:
            if payload.get("path"):
                # API ingestion expects one persisted document id per submission.
                return self._adapter.persist_indexed_document(client, payload)
        except (ImportError, ModuleNotFoundError):
            pass
        doc_id = str(payload.get("doc_id") or payload.get("id") or "").strip()
        if doc_id and doc_id in client.documents:
            return doc_id
        raise RuntimeError("persist failed")

    def _safe_get_pdf_pages(self, file_path: Path) -> int | None:
        """安全获取 PDF 页数，失败时返回空。"""
        try:
            return int(self._adapter.get_pdf_page_count(file_path))
        except Exception:
            return None

    def _safe_get_docx_pages(self, file_path: Path) -> int | None:
        """安全获取 DOCX 逻辑页数，失败时返回空。"""
        try:
            return int(self._adapter.get_docx_page_count(file_path))
        except Exception:
            return None

    def _safe_get_document_pages(self, file_path: Path) -> int | None:
        """按文件类型安全获取页数。

        使用统一多格式解析器获取逻辑页数；失败时返回空。
        """
        try:
            return int(self._adapter.get_document_page_count(file_path))
        except Exception:
            return None

    @staticmethod
    def _estimate_total_nodes(*, total_pages: int | None) -> int | None:
        """基于页数估算最终节点数，用于 structuring 阶段实时展示。"""
        if total_pages is None:
            return None
        pages = max(1, int(total_pages))
        # 在示例文档中节点数通常略高于页数，这里使用保守估算。
        return max(1, int(round(pages * 1.5)))

    @staticmethod
    def _derive_structuring_metrics(
        *,
        progress: int,
        start_progress: int,
        max_progress: int,
        total_pages: int | None,
        estimated_total_nodes: int | None,
    ) -> tuple[int | None, int | None]:
        """根据 structuring 进度推算 processed_pages/tree_node_count。"""
        p = max(int(start_progress), min(int(max_progress), int(progress)))
        span = max(1, int(max_progress) - int(start_progress))
        ratio = max(0.0, min(1.0, float(p - int(start_progress)) / float(span)))

        processed_pages: int | None
        tree_nodes: int | None

        if total_pages is not None:
            processed_pages = max(1, int(round(max(1, int(total_pages)) * ratio)))
        else:
            # 无总页数时仍提供单调递增值，避免长期显示 0。
            processed_pages = max(1, int(round(ratio * 20)))

        if estimated_total_nodes is not None:
            tree_nodes = max(1, int(round(max(1, int(estimated_total_nodes)) * ratio)))
        else:
            tree_nodes = max(1, int(round(ratio * 30)))

        return processed_pages, tree_nodes

    async def _resolve_input_files(self, docs: list[IngestDocument], *, temp_dir: Path) -> tuple[list[Path], list[dict[str, str]]]:
        """把请求中的文档列表解析为可读本地文件路径，并转换为可索引格式。"""
        resolved: list[Path] = []
        source_metadata: list[dict[str, str]] = []
        for doc in docs:
            raw_path = await self._resolve_input_file(
                getattr(doc, "_local_path", None) or "",
                doc.file_name,
                temp_dir=temp_dir,
            )
            prepared = await asyncio.to_thread(
                self._prepare_index_input,
                raw_path,
                declared_file_type=doc.file_type,
                file_name=doc.file_name,
                temp_dir=temp_dir,
            )
            resolved.append(prepared)
            source_metadata.append(
                {
                    "doc_name": str(doc.file_name or raw_path.name),
                    "doc_type": _normalize_storage_doc_type(
                        doc.file_type or raw_path.suffix.lower().lstrip(".")
                    ),
                    "path": str(raw_path),
                    "file_oss_key": str(doc.oss_key or ""),
                }
            )
        return resolved, source_metadata

    def _prepare_index_input(
        self,
        path: Path,
        *,
        declared_file_type: str,
        file_name: str | None = None,
        temp_dir: Path | None = None,
    ) -> Path:
        """将输入文件规范化为可索引格式。

        - 直接支持：pdf/doc/docx/md/markdown/txt/xlsx/pptx
        - html：先转临时 markdown，再复用主索引链路
        """
        suffix = path.suffix.lower()
        if suffix in {".pdf", ".doc", ".docx", ".md", ".markdown", ".txt", ".xlsx", ".pptx"}:
            return path
        if suffix in {".txt", ".html", ".htm"}:
            return self._convert_text_like_to_markdown(path, declared_file_type=declared_file_type, temp_dir=temp_dir)
        declared = str(declared_file_type or "").strip().lower()
        if declared in {"pdf", "doc", "docx", "md", "markdown", "txt", "xlsx", "pptx"}:
            if declared in {"md", "markdown"} and suffix not in {".md", ".markdown"}:
                raise ValueError(
                    f"declared markdown but file suffix is {suffix or 'n/a'}: {path.name}"
                )
            if declared == "txt" and suffix != ".txt":
                raise ValueError(
                    f"declared txt but file suffix is {suffix or 'n/a'}: {path.name}"
                )
            if declared in {"doc", "docx"} and suffix not in {".doc", ".docx"}:
                raise ValueError(
                    f"declared {declared} but file suffix is {suffix or 'n/a'}: {path.name}"
                )
            if declared == "pdf" and suffix != ".pdf":
                raise ValueError(
                    f"declared pdf but file suffix is {suffix or 'n/a'}: {path.name}"
                )
            if declared == "xlsx" and suffix != ".xlsx":
                raise ValueError(
                    f"declared xlsx but file suffix is {suffix or 'n/a'}: {path.name}"
                )
            if declared == "pptx" and suffix != ".pptx":
                raise ValueError(
                    f"declared pptx but file suffix is {suffix or 'n/a'}: {path.name}"
                )
            return path
        if declared == "html":
            return self._convert_text_like_to_markdown(path, declared_file_type=declared, temp_dir=temp_dir)
        raise ValueError(
            f"unsupported input file for indexing: {path.name} (suffix={suffix or 'n/a'}, declared={declared or 'n/a'})"
        )

    @staticmethod
    def _convert_text_like_to_markdown(
        path: Path,
        *,
        declared_file_type: str,
        temp_dir: Path | None = None,
    ) -> Path:
        """将 txt/html 规范化为临时 markdown，复用统一解析链路。"""
        source = path.expanduser().resolve()
        if not source.exists() or not source.is_file():
            raise FileNotFoundError(f"source not found: {source}")
        suffix = source.suffix.lower()
        declared = str(declared_file_type or "").strip().lower()
        is_html = suffix in {".html", ".htm"} or declared == "html"
        try:
            raw = source.read_text(encoding="utf-8", errors="replace")
        except Exception:
            raw = source.read_text(encoding="gb18030", errors="replace")
        if is_html:
            # 保持转换过程尽量简单且可预测：只剥离标签、script 和 style。
            text = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
            text = re.sub(r"(?is)<br\s*/?>", "\n", text)
            text = re.sub(r"(?is)</p\s*>", "\n\n", text)
            text = re.sub(r"(?is)<[^>]+>", " ", text)
            text = re.sub(r"[ \t]+", " ", text)
            text = re.sub(r"\n{3,}", "\n\n", text).strip()
        else:
            text = raw
        tmp_root = Path(temp_dir).resolve() if temp_dir is not None else Path(tempfile.gettempdir()).resolve()
        tmp_root.mkdir(parents=True, exist_ok=True)
        tmp_target = tmp_root / f"ingest_text_{uuid.uuid4().hex}.md"
        tmp_target.write_text(text, encoding="utf-8", errors="replace")
        return tmp_target

    async def _resolve_input_file(self, local_path: str, file_name: str, *, temp_dir: Path) -> Path:
        """解析提交阶段已下载好的本地临时文件。"""
        raw = str(local_path or "").strip()
        if not raw:
            raise ValueError("resolved local file path is required")
        return self._resolve_local_path(raw, file_name=file_name)

    def _resolve_local_path(self, raw: str, *, file_name: str | None = None) -> Path:
        """解析并校验本地文件路径。"""
        p = Path(unquote(raw)).expanduser().resolve()
        if not p.exists() or not p.is_file():
            recovered = self._recover_by_file_name(file_name)
            if recovered is not None:
                self._logger.warning(
                    "resolved file not found; recovered by file_name. requested=%s recovered=%s",
                    p,
                    recovered,
                )
                return recovered
            raise FileNotFoundError(f"resolved file not found: {p}")
        return p

    def _resolve_file_uri(self, parsed, *, file_name: str | None = None) -> Path:
        """解析 file:// URI 并转换为本地路径。"""
        decoded_path = unquote(parsed.path or "")
        if parsed.netloc:
            decoded_path = f"//{parsed.netloc}{decoded_path}"
        # On Windows, file:///C:/path arrives as /C:/path and must drop leading slash.
        if os.name == "nt" and re.match(r"^/[a-zA-Z]:/", decoded_path):
            decoded_path = decoded_path[1:]
        p = Path(decoded_path).expanduser().resolve()
        if not p.exists() or not p.is_file():
            recovered = self._recover_by_file_name(file_name)
            if recovered is not None:
                self._logger.warning(
                    "file_uri not found; recovered by file_name. requested=%s recovered=%s",
                    p,
                    recovered,
                )
                return recovered
            raise FileNotFoundError(f"file URI not found: {p}")
        return p

    def _recover_by_file_name(self, file_name: str | None) -> Path | None:
        """当临时路径失效时，按文件名在候选目录中尝试恢复。"""
        base_name = Path(str(file_name or "").strip()).name
        if not base_name:
            return None

        search_roots: list[Path] = []
        configured = str(os.getenv("INGEST_PATH_SEARCH_ROOTS", "") or "").strip()
        if configured:
            for token in configured.split(os.pathsep):
                candidate = Path(token).expanduser()
                with contextlib.suppress(Exception):
                    search_roots.append(candidate.resolve())

        workspace_path = Path(self.workspace).expanduser()
        with contextlib.suppress(Exception):
            workspace_path = workspace_path.resolve()
        if DEFAULT_SEARCH_ROOTS_ENV:
            default_roots = list(DEFAULT_SEARCH_ROOTS_ENV)
        else:
            default_roots = [
                Path.cwd(),
                _REPO_ROOT,
                _REPO_ROOT / "app",
                _REPO_ROOT / "data",
                workspace_path,
                workspace_path.parent / "data",
                Path("C:/ingest") if os.name == "nt" else Path("/tmp"),
            ]

        seen: set[str] = set()
        ordered_roots: list[Path] = []
        for root in [*search_roots, *default_roots]:
            with contextlib.suppress(Exception):
                resolved = root.resolve()
                key = str(resolved).lower() if os.name == "nt" else str(resolved)
                if key in seen or not resolved.exists() or not resolved.is_dir():
                    continue
                seen.add(key)
                ordered_roots.append(resolved)

        for root in ordered_roots:
            direct_candidates = [
                root / base_name,
                root / "data" / base_name,
            ]
            for candidate in direct_candidates:
                if candidate.exists() and candidate.is_file():
                    return candidate.resolve()

        for root in ordered_roots:
            with contextlib.suppress(Exception):
                for candidate in root.rglob(base_name):
                    if candidate.exists() and candidate.is_file():
                        return candidate.resolve()
        return None

    @staticmethod
    def _download_file(url: str, file_name: str, temp_dir: Path | None = None) -> Path:
        """下载远程文件到临时目录并返回本地路径。"""
        suffix = Path(file_name or "").suffix or ".bin"
        tmp_name = f"ingest_{uuid.uuid4().hex}{suffix}"
        tmp_root = Path(temp_dir).resolve() if temp_dir is not None else Path(tempfile.gettempdir()).resolve()
        tmp_root.mkdir(parents=True, exist_ok=True)
        target = tmp_root / tmp_name
        with requests.get(url, timeout=(5, 20), stream=True) as resp:
            resp.raise_for_status()
            with open(target, "wb") as f:
                for chunk in resp.iter_content(chunk_size=1024 * 1024):
                    if chunk:
                        f.write(chunk)
        return target

    @staticmethod
    def _is_temp_file(path: Path) -> bool:
        """判断文件是否位于系统临时目录。"""
        with contextlib.suppress(Exception):
            tmp_root = Path(tempfile.gettempdir()).resolve()
            return tmp_root in path.resolve().parents
        return False



class ConflictError(Exception):
    """业务冲突异常（例如重复提交中的任务）。"""
    pass


class NotFoundError(Exception):
    """资源不存在异常。"""
    pass


class ServiceUnavailableError(Exception):
    """服务不可用异常（模型加载中/资源不足等）。"""
    pass


class IllegalStatusTransitionError(Exception):
    """状态机非法迁移异常。"""
    pass


class RawMineruRepairError(Exception):
    """Stable internal error contract for on-demand raw repair."""

    def __init__(
        self,
        message: str,
        *,
        code: str,
        retryable: bool,
        http_status: int,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.retryable = retryable
        self.http_status = http_status


@dataclass
class _QueueItem:
    """调度队列中的任务单元。"""
    task_id: str
    request: IngestRequest


class IngestionService:
    """入库任务服务：负责任务提交、调度、查询与删除。"""
    def __init__(
        self,
        store: Any,
        runner: IngestionAlgorithmRunner,
        *,
        callback_dispatcher: CallbackDispatcher | None = None,
        worker_concurrency: int | None = None,
        max_high_burst: int = 2,
        cleanup_hook: Callable[[str, list[tuple[str, str]] | None], Awaitable[int | None] | int | None] | None = None,
    ) -> None:
        """初始化对象并准备运行所需的配置与状态。"""
        self._logger = logging.getLogger(__name__)
        self._expose_runtime_error = str(os.getenv("INGEST_EXPOSE_RUNTIME_ERROR", "") or "").strip().lower() in {
            "1",
            "true",
            "yes",
            "on",
        }
        self._store = store
        self._runner = runner
        self._callback = callback_dispatcher or CallbackDispatcher()
        self._cleanup_hook = cleanup_hook
        self._worker_concurrency = max(1, int(worker_concurrency or os.getenv("INGEST_WORKER_CONCURRENCY", "4")))
        self._temp_worker_concurrency = max(0, int(os.getenv("TEMP_INGEST_WORKER_CONCURRENCY", "2")))
        self._max_high_burst = max(1, int(max_high_burst))
        self._task_timeout_seconds = max(0, int(os.getenv("INGEST_TASK_TIMEOUT_SECONDS", "900")))
        self._high_queue: asyncio.Queue[_QueueItem] = asyncio.Queue()
        self._normal_queue: asyncio.Queue[_QueueItem] = asyncio.Queue()
        self._workers: list[asyncio.Task] = []
        self._worker_lock = asyncio.Lock()
        self._inflight: dict[str, asyncio.Task] = {}
        self._inflight_by_doc: dict[str, set[str]] = {}
        self._cancelled_task_ids: set[str] = set()

    async def _ensure_workers(self) -> None:
        """按并发配置确保 worker 协程已启动。"""
        async with self._worker_lock:
            if self._workers:
                return
            for i in range(self._worker_concurrency):
                self._workers.append(asyncio.create_task(self._worker_loop()))
            for i in range(self._temp_worker_concurrency):
                self._workers.append(asyncio.create_task(self._temp_worker_loop()))

    async def submit_ingestion(self, request: IngestRequest) -> IngestSubmitData:
        """提交普通入库任务并进入常规队列。"""
        return await self._submit_common(
            request=request,
            in_progress_conflict_message="该文档正在处理中，请勿重复提交",
            current_step="queued for processing",
            eta_seconds=self._estimate_ingestion_seconds(request),
            is_temp=False,
            session_id=None,
            high_priority=False,
        )

    async def submit_temp_ingestion(self, request: TempIngestRequest) -> IngestSubmitData:
        """提交临时入库任务并进入高优先级队列。"""
        ingest_request = IngestRequest(
            user_id=request.user_id,
            kb_id=request.kb_id,
            oss_key=request.oss_key,
            file_name=request.file_name,
            file_type=request.file_type,
            callback_url=request.callback_url,
            idempotency_key=request.idempotency_key,
            is_temp=True,
            session_id=request.session_id,
            config=request.config,
        )
        return await self._submit_common(
            request=ingest_request,
            in_progress_conflict_message="该临时文档正在处理中，请勿重复提交",
            current_step="queued in high-priority lane",
            eta_seconds=self._estimate_ingestion_seconds(ingest_request),
            is_temp=True,
            session_id=request.session_id,
            high_priority=True,
        )

    async def submit_raw_mineru_repair(
        self,
        *,
        doc_id: str,
        user_id: str,
        kb_id: str,
    ) -> dict[str, Any]:
        """Claim and enqueue a MinerU-only repair for an existing document."""
        repair_id = f"repair_{uuid.uuid4().hex[:16]}"
        claim = await asyncio.to_thread(
            self._store.claim_raw_mineru_repair,
            doc_id=doc_id,
            user_id=user_id,
            kb_id=kb_id,
            repair_id=repair_id,
        )
        outcome = str(claim.get("outcome") or "")
        repair = dict(claim.get("repair") or {})

        if outcome == "not_found":
            raise RawMineruRepairError(
                "document not found in the requested user_id/kb_id scope",
                code="RAW_MINERU_DOCUMENT_NOT_FOUND",
                retryable=False,
                http_status=404,
            )
        if outcome == "missing_source":
            raise RawMineruRepairError(
                "文档缺少 file_oss_key，无法补充 MinerU 原始结果。请先调用 "
                "DELETE /ingestion/v1/documents/{doc_id} 删除旧文档，再调用标准入库接口重新入库。",
                code="RAW_MINERU_SOURCE_MISSING",
                retryable=False,
                http_status=422,
            )
        if outcome == "terminal_failed":
            raise RawMineruRepairError(
                str(repair.get("error_message") or "MinerU 无法解析该文档，请勿重复重试"),
                code=str(repair.get("error_code") or "RAW_MINERU_REPAIR_TERMINAL"),
                retryable=False,
                http_status=422,
            )
        if outcome in {"complete", "active", "cooldown"}:
            return {
                "doc_id": doc_id,
                "repair_id": repair.get("repair_id"),
                "task_id": repair.get("task_id"),
                "status": "completed" if outcome == "complete" else repair.get("status", "queued"),
                "retryable": bool(repair.get("retryable", outcome != "complete")),
                "next_retry_at": repair.get("next_retry_at"),
                "error_code": repair.get("error_code"),
                "error_message": repair.get("error_message"),
            }
        if outcome != "claimed":
            raise RawMineruRepairError(
                "unable to claim MinerU raw repair",
                code="RAW_MINERU_REPAIR_UNAVAILABLE",
                retryable=True,
                http_status=503,
            )

        document = dict(claim["document"])
        request = IngestRequest(
            user_id=user_id,
            kb_id=kb_id,
            oss_key=document["file_oss_key"],
            file_name=document["doc_name"],
            file_type=document["doc_type"],
        )
        try:
            resolved_doc_ids = await self._resolve_submit_doc_ids(request)
            if resolved_doc_ids != [doc_id]:
                raise RawMineruRepairError(
                    "OSS source content no longer matches the stored document ID; delete and re-ingest the document",
                    code="RAW_MINERU_SOURCE_MISMATCH",
                    retryable=False,
                    http_status=422,
                )
        except Exception as exc:
            self._cleanup_submit_temp_dir(request)
            error = exc if isinstance(exc, RawMineruRepairError) else RawMineruRepairError(
                f"failed to load OSS source: {exc}",
                code="RAW_MINERU_SOURCE_UNAVAILABLE",
                retryable=True,
                http_status=503,
            )
            await asyncio.to_thread(
                self._store.update_raw_mineru_repair,
                doc_id=doc_id,
                repair_id=repair_id,
                updates={
                    "status": "failed",
                    "error_code": error.code,
                    "error_message": str(error),
                    "retryable": error.retryable,
                    "lease_expires_at": None,
                    "next_retry_at": (datetime.now(timezone.utc) + timedelta(seconds=60)).isoformat()
                    if error.retryable
                    else None,
                },
            )
            raise error

        request._doc_id = doc_id
        request._doc_ids = [doc_id]
        request._raw_mineru_repair_doc_ids = {doc_id}
        request._raw_mineru_repair_id = repair_id
        task = TaskStatusData(
            doc_id=doc_id,
            user_id=user_id,
            kb_id=kb_id,
            task_id=f"task_{uuid.uuid4().hex[:16]}",
            status=TaskStatus.QUEUED,
            progress=0,
            current_step="queued for MinerU raw repair",
            updated_at=utc_now_iso(),
            started_at=utc_now_iso(),
            eta_seconds=60,
            persisted_doc_id=doc_id,
            persisted_doc_ids=[doc_id],
            total_documents=1,
        )
        await self._store.upsert(task)
        await asyncio.to_thread(
            self._store.update_raw_mineru_repair,
            doc_id=doc_id,
            repair_id=repair_id,
            updates={"task_id": task.task_id, "status": "queued"},
        )
        await self._ensure_workers()
        await self._normal_queue.put(_QueueItem(task_id=task.task_id, request=request))
        return {
            "doc_id": doc_id,
            "repair_id": repair_id,
            "task_id": task.task_id,
            "status": "queued",
            "retryable": True,
        }

    async def get_raw_mineru_repair_status(
        self,
        *,
        doc_id: str,
        user_id: str,
        kb_id: str,
    ) -> dict[str, Any]:
        """Read repair status without triggering parsing or retries."""
        repair = await asyncio.to_thread(
            self._store.get_raw_mineru_repair,
            doc_id=doc_id,
            user_id=user_id,
            kb_id=kb_id,
        )
        if repair is None:
            raise RawMineruRepairError(
                "document not found in the requested user_id/kb_id scope",
                code="RAW_MINERU_DOCUMENT_NOT_FOUND",
                retryable=False,
                http_status=404,
            )
        return {"doc_id": doc_id, **repair}

    async def _submit_common(
        self,
        *,
        request: IngestRequest,
        in_progress_conflict_message: str,
        current_step: str,
        eta_seconds: int,
        is_temp: bool,
        session_id: str | None,
        high_priority: bool,
    ) -> IngestSubmitData:
        """提交任务公共流程：查重、入队并返回提交结果。"""
        resolved_doc_ids = await self._resolve_submit_doc_ids(request)
        request._doc_ids = resolved_doc_ids
        request._doc_id = resolved_doc_ids[0]

        duplicate_doc_name: str | None = None
        if hasattr(self._store, "get_binding_doc_name"):
            for resolved_doc_id in resolved_doc_ids:
                duplicate_doc_name = await asyncio.to_thread(
                    self._store.get_binding_doc_name,
                    resolved_doc_id,
                    request.user_id,
                    request.kb_id,
                )
                if duplicate_doc_name:
                    break
        if duplicate_doc_name:
            if is_temp:
                existing = await self._find_existing_task(
                    doc_id=request.doc_id,
                    user_id=request.user_id,
                    kb_id=request.kb_id,
                    idempotency_key=request.idempotency_key,
                )
                if existing is None:
                    now = utc_now_iso()
                    existing = TaskStatusData(
                        doc_id=request.doc_id,
                        user_id=request.user_id,
                        kb_id=request.kb_id,
                        task_id=f"reuse_{uuid.uuid4().hex[:16]}",
                        status=TaskStatus.COMPLETED,
                        progress=100,
                        current_step="reused existing document",
                        updated_at=now,
                        started_at=now,
                        completed_at=now,
                        eta_seconds=0,
                        persisted_doc_id=request.doc_id,
                        persisted_doc_ids=[request.doc_id],
                        total_documents=1,
                        processed_documents=1,
                        callback_url=str(request.callback_url) if request.callback_url else None,
                        idempotency_key=request.idempotency_key,
                        is_temp=True,
                        session_id=session_id,
                    )
                    await self._store.upsert(existing)
                self._cleanup_submit_temp_dir(request)
                return IngestSubmitData(
                    task_id=existing.task_id,
                    doc_id=request.doc_id,
                    doc_ids=[request.doc_id],
                    user_id=request.user_id,
                    status="completed",
                    estimated_seconds=0,
                    is_duplicate=True,
                    document_count=1,
                )
            can_repair_raw = (
                not is_temp
                and request.document_count == 1
                and str(os.getenv("INGEST_PARSER_BACKEND", "native")).strip().lower() == "mineru"
                and hasattr(self._store, "has_complete_raw_mineru")
            )
            if can_repair_raw:
                raw_complete = await asyncio.to_thread(
                    self._store.has_complete_raw_mineru,
                    resolved_doc_ids[0],
                )
                if not raw_complete:
                    repair_id = f"repair_{uuid.uuid4().hex[:16]}"
                    claim = await asyncio.to_thread(
                        self._store.claim_raw_mineru_repair,
                        doc_id=resolved_doc_ids[0],
                        user_id=request.user_id,
                        kb_id=request.kb_id,
                        repair_id=repair_id,
                        require_source=False,
                    )
                    outcome = str(claim.get("outcome") or "")
                    repair = dict(claim.get("repair") or {})
                    if outcome == "claimed":
                        request._raw_mineru_repair_doc_ids.add(resolved_doc_ids[0])
                        request._raw_mineru_repair_id = repair_id
                    elif outcome == "active":
                        self._cleanup_submit_temp_dir(request)
                        raise ConflictError("该文档的 MinerU raw 修复正在处理中，请勿重复提交")
                    elif outcome in {"cooldown", "terminal_failed"}:
                        self._cleanup_submit_temp_dir(request)
                        retryable = outcome == "cooldown" or bool(repair.get("retryable"))
                        raise RawMineruRepairError(
                            str(repair.get("error_message") or "MinerU raw repair is unavailable"),
                            code=str(repair.get("error_code") or "RAW_MINERU_REPAIR_FAILED"),
                            retryable=retryable,
                            http_status=503 if retryable else 422,
                        )
                    elif outcome != "complete":
                        self._cleanup_submit_temp_dir(request)
                        raise RawMineruRepairError(
                            "unable to claim MinerU raw repair",
                            code="RAW_MINERU_REPAIR_UNAVAILABLE",
                            retryable=True,
                            http_status=503,
                        )
            if not request.raw_mineru_repair_doc_ids:
                self._cleanup_submit_temp_dir(request)
                raise DuplicateDocumentError(duplicate_doc_name)

        existing = await self._find_existing_task(
            doc_id=request.doc_id,
            user_id=request.user_id,
            kb_id=request.kb_id,
            idempotency_key=request.idempotency_key,
        )
        if existing is not None:
            if existing.status in IN_PROGRESS_STATUSES:
                self._cleanup_submit_temp_dir(request)
                raise ConflictError(in_progress_conflict_message)
            if existing.status == TaskStatus.COMPLETED and not request.raw_mineru_repair_doc_ids:
                self._cleanup_submit_temp_dir(request)
                return IngestSubmitData(
                    task_id=existing.task_id,
                    doc_id=existing.doc_id,
                    doc_ids=existing.persisted_doc_ids or [existing.doc_id],
                    user_id=request.user_id,
                    status="completed",
                    estimated_seconds=0,
                    is_duplicate=True,
                    document_count=request.document_count,
                )

        task_data = TaskStatusData(
            doc_id=request.doc_id,
            user_id=request.user_id,
            kb_id=request.kb_id,
            task_id=f"task_{uuid.uuid4().hex[:16]}",
            status=TaskStatus.QUEUED,
            progress=0,
            current_step=current_step,
            updated_at=utc_now_iso(),
            started_at=utc_now_iso(),
            eta_seconds=eta_seconds,
            persisted_doc_id=request.doc_id,
            persisted_doc_ids=request.doc_ids,
            total_documents=request.document_count,
            processed_documents=0,
            callback_url=str(request.callback_url) if request.callback_url else None,
            idempotency_key=request.idempotency_key,
            is_temp=is_temp,
            session_id=session_id,
        )
        await self._store.upsert(task_data)
        if request.raw_mineru_repair_id and hasattr(self._store, "update_raw_mineru_repair"):
            await asyncio.to_thread(
                self._store.update_raw_mineru_repair,
                doc_id=request.doc_id,
                repair_id=request.raw_mineru_repair_id,
                updates={"task_id": task_data.task_id, "status": "queued"},
            )
        await self._ensure_workers()
        queue = self._high_queue if high_priority else self._normal_queue
        await queue.put(_QueueItem(task_id=task_data.task_id, request=request))
        return IngestSubmitData(
            task_id=task_data.task_id,
            doc_id=request.doc_id,
            doc_ids=request.doc_ids,
            user_id=request.user_id,
            status="queued",
            estimated_seconds=task_data.eta_seconds or 0,
            is_duplicate=False,
            document_count=request.document_count,
        )

    @staticmethod
    def _cleanup_submit_temp_dir(request: IngestRequest) -> None:
        """清理提交阶段为 HTTP 文件下载创建的临时目录。"""
        submit_temp_dir = Path(request._submit_temp_dir).resolve() if request._submit_temp_dir else None
        if submit_temp_dir is None:
            return
        with contextlib.suppress(Exception):
            if submit_temp_dir.exists():
                shutil.rmtree(submit_temp_dir, ignore_errors=True)
        request._submit_temp_dir = None

    async def _resolve_submit_doc_ids(self, request: IngestRequest) -> list[str]:
        """提交阶段先确认每个文档的内容 UUID。

        OSS 文件会先完整下载并计算 sha256，并把下载后的临时文件路径写回 request，
        供后台解析复用。
        """
        docs = request.iter_documents()
        submit_temp_dir: Path | None = None
        resolved: list[str] = []
        for index, doc in enumerate(docs):
            doc_id, local_path = await asyncio.to_thread(
                self._stable_document_id_and_local_path,
                doc.oss_key,
                doc.file_name,
                submit_temp_dir,
            )
            if local_path is not None:
                submit_temp_dir = local_path.parent
                if request.documents:
                    request.documents[index]._local_path = str(local_path)
                else:
                    request._single_local_path = str(local_path)
            resolved.append(doc_id)
        if submit_temp_dir is not None:
            request._submit_temp_dir = str(submit_temp_dir)
        return resolved or [request.doc_id]

    @staticmethod
    def _unique_download_target(temp_dir: Path, file_name: str) -> Path:
        """Build a collision-safe temp path while preserving the display name."""
        raw_name = Path(str(file_name or "").strip()).name
        safe_name = raw_name or f"download_{uuid.uuid4().hex}.bin"
        target = temp_dir / safe_name
        if not target.exists():
            return target

        stem = Path(safe_name).stem or "download"
        suffix = Path(safe_name).suffix
        for index in range(1, 1000):
            candidate = temp_dir / f"{stem}_{index}{suffix}"
            if not candidate.exists():
                return candidate
        return temp_dir / f"{stem}_{uuid.uuid4().hex}{suffix}"

    @staticmethod
    def _stable_document_id_and_local_path(
        oss_key: str,
        file_name: str,
        submit_temp_dir: Path | None = None,
    ) -> tuple[str, Path | None]:
        """根据 OSS 文件内容生成稳定文档 UUID。"""
        source_oss_key = str(oss_key or "").strip()
        if not source_oss_key:
            raise ValueError("oss_key is required")

        sha256 = hashlib.sha256()
        suffix = Path(file_name or "").suffix or ".bin"
        tmp_root = Path(submit_temp_dir).resolve() if submit_temp_dir is not None else Path(
            tempfile.mkdtemp(prefix="ingest_submit_")
        ).resolve()
        tmp_root.mkdir(parents=True, exist_ok=True)
        target = IngestionService._unique_download_target(tmp_root, file_name)
        if not target.suffix:
            target = target.with_suffix(suffix)

        from app.ingestion.oss_client import download_oss_key

        download_oss_key(source_oss_key, target)
        with open(target, "rb") as f:
            for chunk in iter(lambda: f.read(1024 * 1024), b""):
                sha256.update(chunk)
        content_hash = sha256.hexdigest()
        doc_identity = f"sha256:{content_hash}"
        return str(uuid.uuid5(uuid.NAMESPACE_URL, doc_identity)), target

    async def get_status_by_doc_id(self, doc_id: str, *, user_id: str, kb_id: str) -> TaskStatusData:
        """按用户、知识库与落库文档 UUID 获取最新任务状态。"""
        try:
            doc_id = str(uuid.UUID(str(doc_id)))
        except (TypeError, ValueError) as exc:
            raise NotFoundError(f"document uuid not found: {doc_id}") from exc
        task = await self._store.get_latest_by_persisted_doc_id(doc_id, user_id, kb_id)
        if task is None:
            raise NotFoundError(f"user_id/kb_id/document uuid not found: {user_id}/{kb_id}/{doc_id}")
        return task

    async def get_status_by_task_id(self, task_id: str) -> TaskStatusData:
        """按任务 ID 获取任务状态。"""
        task = await self._store.get_by_task_id(task_id)
        if task is None:
            raise NotFoundError(f"task_id not found: {task_id}")
        return task

    async def delete_document(
        self,
        doc_id: str,
        *,
        user_id: str,
        kb_id: str,
        cancel_running: bool = False,
    ) -> dict[str, Any]:
        """删除文档相关任务（硬删除，按 user_id + kb_id + doc_id）。"""
        rows = await self._store.get_all_by_doc_id(doc_id, user_id=user_id, kb_id=kb_id)
        if not rows:
            raise NotFoundError(f"user_id/kb_id/doc_id not found: {user_id}/{kb_id}/{doc_id}")
        # 识别运行中任务，决定是否允许删除。
        running = [r for r in rows if r.status in IN_PROGRESS_STATUSES]
        if running and not cancel_running:
            raise ConflictError("document has running ingestion task; set cancel_running=true")

        if running and cancel_running:
            for r in running:
                self._cancelled_task_ids.add(r.task_id)
                # 先尝试取消协程，再把任务状态写成已取消失败。
                task = self._inflight.get(r.task_id)
                if task and not task.done():
                    task.cancel()
                await self._store.update_status(
                    r.task_id,
                    status=TaskStatus.FAILED,
                    progress=min(99, max(0, int(r.progress))),
                    current_step="task cancelled by delete request",
                    error_message="cancelled by delete",
                    error_code="TASK_CANCELLED",
                    retryable=False,
                    eta_seconds=0,
                )

        # 汇总落库文档 ID，用于后续存储层清理钩子。
        persisted_ids: list[str] = []
        seen_ids: set[str] = set()
        for row in rows:
            single = str(row.persisted_doc_id or "").strip()
            if single and single not in seen_ids:
                seen_ids.add(single)
                persisted_ids.append(single)
            for raw in row.persisted_doc_ids:
                candidate = str(raw or "").strip()
                if candidate and candidate not in seen_ids:
                    seen_ids.add(candidate)
                    persisted_ids.append(candidate)

        # 清理任务记录、取消标记和回调死信残留。
        removed = await self._store.remove_by_doc_id(doc_id, user_id=user_id, kb_id=kb_id)
        deleted_task_ids = {row.task_id for row in rows}
        for row in rows:
            self._cancelled_task_ids.discard(row.task_id)
        self._callback.dlq = [item for item in self._callback.dlq if item.task_id not in deleted_task_ids]
        deleted_tree_nodes = 0
        if self._cleanup_hook is not None:
            targets = persisted_ids or [doc_id]
            bindings = sorted({(str(row.user_id), str(row.kb_id)) for row in rows})
            for target_doc_id in targets:
                maybe = self._cleanup_hook(target_doc_id, bindings)
                if asyncio.iscoroutine(maybe):
                    maybe = await maybe
                if maybe is not None:
                    deleted_tree_nodes += max(0, int(maybe))
        elif removed > 0:
            # Fallback for environments without storage cleanup hook.
            deleted_tree_nodes = int(removed)
        return {"doc_id": doc_id, "deleted_tree_nodes": deleted_tree_nodes}

    async def delete_session_temp_documents(
        self,
        session_id: str,
        *,
        kb_id: str | None = None,
        cancel_running: bool = True,
    ) -> dict[str, Any]:
        """清理某个 session 命名空间下的全部临时文档。"""
        session_id = str(session_id or "").strip()
        if not session_id:
            raise ValueError("session_id is required")
        rows = await self._store.get_temp_by_session(session_id, kb_id=kb_id)
        if not rows:
            return {
                "session_id": session_id,
                "kb_id": kb_id,
                "deleted_temp_docs": 0,
                "deleted_tree_nodes": 0,
                "doc_ids": [],
            }

        seen_doc_keys: set[tuple[str, str, str]] = set()
        targets: list[tuple[str, str, str]] = []
        for row in rows:
            key = (str(row.doc_id), str(row.user_id), str(row.kb_id))
            if key in seen_doc_keys:
                continue
            seen_doc_keys.add(key)
            targets.append(key)

        deleted_tree_nodes = 0
        deleted_doc_ids: list[str] = []
        for doc_id, row_user_id, row_kb_id in targets:
            result = await self.delete_document(
                doc_id,
                user_id=row_user_id,
                kb_id=row_kb_id,
                cancel_running=cancel_running,
            )
            deleted_doc_ids.append(doc_id)
            deleted_tree_nodes += max(0, int(result.get("deleted_tree_nodes") or 0))

        return {
            "session_id": session_id,
            "kb_id": kb_id,
            "deleted_temp_docs": len(deleted_doc_ids),
            "deleted_tree_nodes": deleted_tree_nodes,
            "doc_ids": deleted_doc_ids,
        }

    async def _worker_loop(self) -> None:
        """工作协程主循环：按策略消费高优先级与普通队列。"""
        high_streak = 0
        while True:
            item = None
            # 高优队列支持突发消费，但限制连续次数避免普通队列饥饿。
            if (not self._high_queue.empty()) and (high_streak < self._max_high_burst or self._normal_queue.empty()):
                try:
                    item = self._high_queue.get_nowait()
                    high_streak += 1
                except asyncio.QueueEmpty:
                    item = None
            elif not self._normal_queue.empty():
                try:
                    item = self._normal_queue.get_nowait()
                    high_streak = 0
                except asyncio.QueueEmpty:
                    item = None
            elif not self._high_queue.empty():
                try:
                    item = self._high_queue.get_nowait()
                    high_streak = 1
                except asyncio.QueueEmpty:
                    item = None

            if item is None:
                await asyncio.sleep(0.05)
                continue

            await self._execute_queue_item(item)

    async def _temp_worker_loop(self) -> None:
        """临时文档专用 worker：只消费高优先级队列。"""
        while True:
            try:
                item = await self._high_queue.get()
            except asyncio.CancelledError:
                raise
            await self._execute_queue_item(item)

    async def _execute_queue_item(self, item: _QueueItem) -> None:
        """执行队列项，统一处理取消检查与 inflight 记录。"""
        # 文档已被取消时，直接写入失败终态并跳过执行。
        if item.task_id in self._cancelled_task_ids:
            current = await self._store.get_by_task_id(item.task_id)
            failed_progress = min(99, max(0, int(current.progress if current is not None else 0)))
            await self._store.update_status(
                item.task_id,
                status=TaskStatus.FAILED,
                progress=failed_progress,
                current_step="任务因取消而跳过",
                error_message="cancelled before execution",
                error_code="TASK_CANCELLED",
                retryable=False,
                eta_seconds=0,
            )
            return

        task = asyncio.create_task(self._run_task(item.task_id, item.request))
        self._inflight[item.task_id] = task
        self._inflight_by_doc.setdefault(item.request.doc_id, set()).add(item.task_id)
        with contextlib.suppress(Exception):
            await task

    async def _run_task(self, task_id: str, request: IngestRequest) -> None:
        """执行单个任务并统一处理取消、超时、异常与回调。"""
        async def _failed_progress() -> int:
            current = await self._store.get_by_task_id(task_id)
            base = 0 if current is None else int(current.progress)
            return min(99, max(0, base))

        async def _mark_failed(
            *,
            current_step: str,
            error_message: str,
            error_code: str,
            retryable: bool,
        ) -> None:
            await self._store.update_status(
                task_id,
                status=TaskStatus.FAILED,
                progress=await _failed_progress(),
                current_step=current_step,
                error_message=error_message,
                error_code=error_code,
                retryable=retryable,
                eta_seconds=0,
            )

        async def report(**kwargs) -> None:
            """包装进度上报，附带状态机迁移合法性校验。"""
            status = TaskStatus(kwargs.pop("status"))
            current = await self._store.get_by_task_id(task_id)
            # 防止非法状态迁移污染任务状态轨迹。
            if current is not None and not self._is_valid_transition(current.status, status):
                await _mark_failed(
                    current_step="illegal status transition",
                    error_message=f"illegal transition: {current.status} -> {status}",
                    error_code="ILLEGAL_STATUS_TRANSITION",
                    retryable=False,
                )
                raise IllegalStatusTransitionError("illegal status transition")
            await self._store.update_status(task_id, status=status, **kwargs)

        try:
            # 为整条执行链增加超时保护，避免任务长期悬挂。
            if self._task_timeout_seconds > 0:
                await asyncio.wait_for(self._runner.run(request, report), timeout=self._task_timeout_seconds)
            else:
                await self._runner.run(request, report)
        except asyncio.CancelledError:
            await _mark_failed(
                current_step="task cancelled",
                error_message="cancelled",
                error_code="TASK_CANCELLED",
                retryable=False,
            )
        except asyncio.TimeoutError:
            await _mark_failed(
                current_step="task timed out",
                error_message=f"task exceeded timeout ({self._task_timeout_seconds}s)",
                error_code="INGESTION_TIMEOUT",
                retryable=True,
            )
        except IllegalStatusTransitionError:
            pass
        except RawMineruRepairTerminalError as exc:
            await _mark_failed(
                current_step="MinerU raw repair failed",
                error_message=str(exc),
                error_code="RAW_MINERU_INVALID_RESULT",
                retryable=False,
            )
        except DuplicateDocumentError as exc:
            self._logger.info("ingestion task duplicate document: task_id=%s doc_id=%s", task_id, request.doc_id)
            await _mark_failed(
                current_step="duplicate document",
                error_message=str(exc),
                error_code="DUPLICATE_DOCUMENT",
                retryable=False,
            )
        except Exception as exc:
            self._logger.exception("ingestion task failed: task_id=%s doc_id=%s", task_id, request.doc_id)
            await _mark_failed(
                current_step="task execution failed",
                error_message=str(exc) if self._expose_runtime_error else "runtime error",
                error_code="INGESTION_RUNTIME_ERROR",
                retryable=True,
            )
        finally:
            final = await self._store.get_by_task_id(task_id)
            if (
                request.raw_mineru_repair_id
                and final is not None
                and final.status == TaskStatus.FAILED
                and hasattr(self._store, "update_raw_mineru_repair")
            ):
                retryable = bool(final.retryable)
                now = datetime.now(timezone.utc)
                await asyncio.to_thread(
                    self._store.update_raw_mineru_repair,
                    doc_id=request.doc_id,
                    repair_id=request.raw_mineru_repair_id,
                    updates={
                        "status": "failed",
                        "error_code": final.error_code,
                        "error_message": final.error_message,
                        "retryable": retryable,
                        "updated_at": now.isoformat(),
                        "lease_expires_at": None,
                        "next_retry_at": (now + timedelta(seconds=60)).isoformat() if retryable else None,
                    },
                )
            # 仅在终态触发回调，避免中间态频繁通知。
            if final is not None and final.status in {TaskStatus.COMPLETED, TaskStatus.FAILED}:
                await self._callback.dispatch(final.callback_url, final)
            self._inflight.pop(task_id, None)
            ids = self._inflight_by_doc.get(request.doc_id)
            if ids:
                ids.discard(task_id)
                if not ids:
                    self._inflight_by_doc.pop(request.doc_id, None)

    async def _find_existing_task(
        self,
        *,
        doc_id: str,
        user_id: str,
        kb_id: str,
        idempotency_key: str | None,
    ) -> TaskStatusData | None:
        """根据幂等键或业务键查找已有任务。"""
        if idempotency_key and hasattr(self._store, "get_by_idempotency_key"):
            hit = await self._store.get_by_idempotency_key(kb_id, user_id, idempotency_key)
            if hit is not None:
                return hit
        return await self._store.get_by_biz_key(doc_id, user_id, kb_id)

    @staticmethod
    def _is_valid_transition(previous: TaskStatus, nxt: TaskStatus) -> bool:
        """校验状态机迁移是否合法。"""
        if previous == nxt:
            return True
        allowed = {
            TaskStatus.QUEUED: {TaskStatus.PARSING, TaskStatus.FAILED},
            TaskStatus.PARSING: {TaskStatus.EXTRACTING, TaskStatus.FAILED},
            TaskStatus.EXTRACTING: {TaskStatus.STRUCTURING, TaskStatus.FAILED},
            TaskStatus.STRUCTURING: {TaskStatus.STORING, TaskStatus.FAILED},
            TaskStatus.STORING: {TaskStatus.COMPLETED, TaskStatus.FAILED},
            TaskStatus.COMPLETED: set(),
            TaskStatus.FAILED: set(),
        }
        return nxt in allowed.get(previous, set())

    @staticmethod
    def _estimate_ingestion_seconds(request: IngestRequest) -> int:
        """按并发模型粗估任务耗时，用于前端展示 ETA。

        这里估算的是“批次耗时”而不是“页数 * 单页耗时”：
        - 视觉 OCR/复杂页解析会按 pdf_vision_concurrency 并发；
        - 摘要生成会按 summary_concurrency 并发；
        - 普通解析与存储按轻量固定成本估算。
        """
        vision_page_avg_seconds = float(os.getenv("INGEST_ETA_VISION_PAGE_SECONDS", "8"))
        summary_node_avg_seconds = float(os.getenv("INGEST_ETA_SUMMARY_NODE_SECONDS", "0.8"))
        normal_parse_page_avg_seconds = float(os.getenv("INGEST_ETA_NORMAL_PARSE_PAGE_SECONDS", "0.05"))
        store_avg_seconds = float(os.getenv("INGEST_ETA_STORE_SECONDS", "2"))

        vision_concurrency = max(
            1,
            int(
                os.getenv("PDF_VISION_CONCURRENCY")
                or os.getenv("DA_PDF_VISION_CONCURRENCY")
                or os.getenv("ARK_VISION_CONCURRENCY")
                or "32"
            ),
        )
        summary_concurrency = max(
            1,
            int(
                os.getenv("SUMMARY_CONCURRENCY")
                or os.getenv("DA_SUMMARY_CONCURRENCY")
                or os.getenv("ARK_SUMMARY_CONCURRENCY")
                or "32"
            ),
        )

        total_seconds = 0.0
        for item in request.iter_documents():
            file_type = str(item.file_type or "").lower()
            page_count = IngestionService._estimate_document_page_count(item)
            estimated_nodes = IngestionService._estimate_summary_node_count(file_type, page_count)

            parse_seconds = page_count * normal_parse_page_avg_seconds
            # 目前只有 PDF 的复杂页/OCR 页会走 Ark Vision；提交阶段无法完成复杂页扫描，
            # 所以在 enable_ocr=true 时按“可能需要视觉解析的页数”做保守估计。
            vision_pages = page_count if file_type == "pdf" and bool(request.config.enable_ocr) else 0
            vision_seconds = math.ceil(vision_pages / vision_concurrency) * vision_page_avg_seconds if vision_pages else 0.0
            summary_seconds = (
                math.ceil(estimated_nodes / summary_concurrency) * summary_node_avg_seconds
                if bool(request.config.summary_enabled) and estimated_nodes
                else 0.0
            )
            total_seconds += parse_seconds + vision_seconds + summary_seconds + store_avg_seconds

        return max(3, int(math.ceil(total_seconds)))

    @staticmethod
    def _estimate_summary_node_count(file_type: str, page_count: int) -> int:
        """根据文档类型粗估需要生成摘要的节点数。"""
        pages = max(1, int(page_count or 1))
        if file_type == "pdf":
            return max(1, math.ceil(pages / 2))
        if file_type == "pptx":
            return pages
        if file_type == "xlsx":
            return pages
        return max(1, math.ceil(pages / 3))

    @staticmethod
    def _estimate_document_page_count(item: IngestDocument) -> int:
        """尽量快速估算页数；读不到真实页数时按文件大小粗估。"""
        file_type = str(item.file_type or "").lower()
        type_default_pages = {
            "pdf": 10,
            "doc": 5,
            "pptx": 10,
            "xlsx": 3,
            "docx": 5,
            "txt": 3,
            "md": 3,
            "markdown": 3,
            "html": 3,
        }.get(file_type, 3)
        local_path = getattr(item, "_local_path", None)
        size_bytes = IngestionService._estimate_file_size_bytes(local_path)
        default_pages = IngestionService._estimate_page_count_by_size(
            file_type=file_type,
            size_bytes=size_bytes,
            fallback=type_default_pages,
        )

        path = IngestionService._local_path_from_url(local_path)
        if path is None or not path.exists() or not path.is_file():
            return default_pages

        try:
            if file_type == "pdf":
                try:
                    import fitz  # type: ignore

                    with fitz.open(str(path)) as doc:
                        return max(1, int(doc.page_count))
                except Exception:
                    return default_pages
            if file_type == "pptx":
                try:
                    from pptx import Presentation  # type: ignore

                    return max(1, len(Presentation(str(path)).slides))
                except Exception:
                    return default_pages
            if file_type == "xlsx":
                try:
                    from openpyxl import load_workbook  # type: ignore

                    wb = load_workbook(str(path), read_only=True, data_only=True)
                    try:
                        return max(1, len(wb.sheetnames))
                    finally:
                        wb.close()
                except Exception:
                    return default_pages
            if file_type in {"txt", "md", "markdown", "html"}:
                size = max(1, path.stat().st_size)
                # 粗略按 UTF-8 中文/英文混排 2 bytes/char、1500 chars/node 估算逻辑页。
                return max(1, math.ceil(size / 3000))
        except Exception:
            return default_pages
        return default_pages

    @staticmethod
    def _estimate_page_count_by_size(*, file_type: str, size_bytes: int | None, fallback: int) -> int:
        """用文件大小粗估逻辑页数，避免提交阶段所有 PDF 都落到固定默认值。"""
        if not size_bytes or int(size_bytes) <= 0:
            return max(1, int(fallback))

        # 这些值只影响提交响应里的 estimated_seconds，不影响真实解析结果。
        bytes_per_page = {
            "pdf": int(os.getenv("INGEST_ETA_PDF_BYTES_PER_PAGE", "200000")),
            "pptx": int(os.getenv("INGEST_ETA_PPTX_BYTES_PER_PAGE", "400000")),
            "xlsx": int(os.getenv("INGEST_ETA_XLSX_BYTES_PER_PAGE", "120000")),
            "doc": int(os.getenv("INGEST_ETA_DOCX_BYTES_PER_PAGE", "80000")),
            "docx": int(os.getenv("INGEST_ETA_DOCX_BYTES_PER_PAGE", "80000")),
            "txt": int(os.getenv("INGEST_ETA_TEXT_BYTES_PER_PAGE", "3000")),
            "md": int(os.getenv("INGEST_ETA_TEXT_BYTES_PER_PAGE", "3000")),
            "markdown": int(os.getenv("INGEST_ETA_TEXT_BYTES_PER_PAGE", "3000")),
            "html": int(os.getenv("INGEST_ETA_TEXT_BYTES_PER_PAGE", "3000")),
        }.get(file_type, 120000)
        bytes_per_page = max(1, int(bytes_per_page))
        return max(1, int(math.ceil(int(size_bytes) / bytes_per_page)))

    @staticmethod
    def _estimate_file_size_bytes(source: str | None) -> int | None:
        """尽量获取文件大小；本地文件用 stat，http(s) 用 HEAD 的 Content-Length。"""
        path = IngestionService._local_path_from_url(source)
        if path is not None and path.exists() and path.is_file():
            with contextlib.suppress(OSError):
                return int(path.stat().st_size)

        value = str(source or "").strip()
        parsed = urlparse(value)
        if parsed.scheme not in {"http", "https"}:
            return None
        try:
            resp = requests.head(value, allow_redirects=True, timeout=3)
            content_length = resp.headers.get("content-length")
            if content_length:
                return int(content_length)
        except Exception:
            return None
        return None

    @staticmethod
    def _local_path_from_url(source: str | None) -> Path | None:
        """将本地路径或 file:// URL 转为 Path；http(s) 返回 None。"""
        value = str(source or "").strip()
        if not value:
            return None
        if re.match(r"^[A-Za-z]:[\\/]", value):
            return Path(value)
        parsed = urlparse(value)
        if parsed.scheme in {"http", "https"}:
            return None
        if len(parsed.scheme) == 1 and re.match(r"^[A-Za-z]$", parsed.scheme):
            # urlparse("C:/path/file.pdf") 会把盘符识别为 scheme="c"。
            return Path(value)
        if parsed.scheme == "file":
            raw_path = unquote(parsed.path or "")
            if os.name == "nt" and re.match(r"^/[A-Za-z]:/", raw_path):
                raw_path = raw_path[1:]
            return Path(raw_path)
        if parsed.scheme:
            return None
        return Path(value)

try:
    from fastapi import APIRouter, Depends, FastAPI, HTTPException, Request
    from fastapi.encoders import jsonable_encoder
    from fastapi.exceptions import RequestValidationError
    from fastapi.responses import JSONResponse

    _FASTAPI_AVAILABLE = True
except Exception:  # pragma: no cover - fastapi may be absent in minimal runtime
    APIRouter = Depends = FastAPI = HTTPException = Request = None  # type: ignore[assignment]
    RequestValidationError = Exception  # type: ignore[assignment]
    JSONResponse = None  # type: ignore[assignment]
    _FASTAPI_AVAILABLE = False


_SERVICE: IngestionService | None = None


if _FASTAPI_AVAILABLE:
    class Utf8JSONResponse(JSONResponse):
        """统一声明 UTF-8，避免部分客户端出现中文乱码。"""

        media_type = "application/json; charset=utf-8"
else:  # pragma: no cover - FastAPI 不可用时的回退路径
    Utf8JSONResponse = None  # type: ignore[assignment]


def get_ingestion_service() -> IngestionService:
    """FastAPI 依赖注入入口：返回全局 IngestionService 实例。"""
    if _SERVICE is None:
        raise ServiceUnavailableError("ingestion service is not initialized")
    return _SERVICE


if _FASTAPI_AVAILABLE:
    router = APIRouter(prefix=API_PREFIX, tags=["ingestion"])

    @router.post("/documents/ingest", response_model=ApiResponse)
    async def ingest_document(
        request: IngestRequest,
        service: IngestionService = Depends(get_ingestion_service),
    ) -> ApiResponse:
        """普通文档入库 API 接口。"""
        try:
            data = await service.submit_ingestion(request)
        except ConflictError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        except DuplicateDocumentError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        except (FileNotFoundError, RuntimeError, requests.RequestException, ValueError) as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        payload = {
            "task_id": data.task_id,
            "doc_id": data.doc_id,
            "doc_ids": data.doc_ids,
            "user_id": data.user_id,
            "status": data.status,
            "estimated_seconds": data.estimated_seconds,
        }
        return ApiResponse(
            code=200,
            message="文档已提交处理",
            data=payload,
        )

    @router.post(
        "/internal/documents/raw/repair",
        response_model=ApiResponse,
        status_code=202,
    )
    async def repair_raw_mineru(
        request: RawMineruRepairRequest,
        service: IngestionService = Depends(get_ingestion_service),
    ) -> ApiResponse:
        """Queue an in-place MinerU repair for an existing document."""
        data = await service.submit_raw_mineru_repair(
            doc_id=request.doc_id,
            user_id=request.user_id,
            kb_id=request.kb_id,
        )
        return ApiResponse(code=202, message="MinerU raw repair accepted", data=data)

    @router.post("/documents/temp-ingest", response_model=ApiResponse)
    async def temp_ingest_document(
        request: TempIngestRequest,
        service: IngestionService = Depends(get_ingestion_service),
    ) -> ApiResponse:
        """临时文档入库 API 接口。"""
        try:
            data = await service.submit_temp_ingestion(request)
        except ConflictError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        except DuplicateDocumentError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        except (FileNotFoundError, RuntimeError, requests.RequestException, ValueError) as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        payload = {
            "task_id": data.task_id,
            "doc_id": data.doc_id,
            "doc_ids": data.doc_ids,
            "user_id": data.user_id,
            "status": data.status,
            "estimated_seconds": data.estimated_seconds,
        }
        return ApiResponse(
            code=200,
            message="临时文档已提交处理（高优先级）",
            data=payload,
        )

    @router.get("/documents/{doc_id}/status", response_model=ApiResponse)
    async def get_status_by_doc(
        doc_id: str,
        user_id: str,
        kb_id: str,
        service: IngestionService = Depends(get_ingestion_service),
    ) -> ApiResponse:
        """按用户、知识库与落库文档 UUID 查询任务状态 API 接口。"""
        try:
            data = await service.get_status_by_doc_id(doc_id, user_id=user_id, kb_id=kb_id)
        except NotFoundError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        return ApiResponse(code=200, message="success", data=_public_status_payload(data))

    @router.get("/tasks/{task_id}", response_model=ApiResponse)
    async def get_status_by_task(
        task_id: str,
        service: IngestionService = Depends(get_ingestion_service),
    ) -> ApiResponse:
        """按任务 ID 查询任务状态 API 接口。"""
        try:
            data = await service.get_status_by_task_id(task_id)
        except NotFoundError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        return ApiResponse(code=200, message="success", data=_public_status_payload(data))

    @router.delete("/documents/{doc_id}", response_model=ApiResponse)
    async def delete_document(
        doc_id: str,
        user_id: str,
        kb_id: str,
        is_temp: bool = False,
        service: IngestionService = Depends(get_ingestion_service),
    ) -> ApiResponse:
        """删除文档任务数据（硬删除）。"""
        try:
            # 删除范围由 user_id + kb_id + doc_id 唯一定位，避免共享内容 UUID 串用户。
            data = await service.delete_document(
                doc_id,
                user_id=user_id,
                kb_id=kb_id,
                cancel_running=False,
            )
        except ConflictError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        except NotFoundError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        return ApiResponse(code=200, message="文档数据已删除", data=data)

    @router.delete("/sessions/{session_id}/temp-documents", response_model=ApiResponse)
    async def delete_session_temp_documents(
        session_id: str,
        kb_id: str | None = None,
        service: IngestionService = Depends(get_ingestion_service),
    ) -> ApiResponse:
        """会话销毁时清理该 session 命名空间下的临时文档。"""
        data = await service.delete_session_temp_documents(
            session_id,
            kb_id=kb_id,
            cancel_running=True,
        )
        return ApiResponse(code=200, message="会话临时文档已删除", data=data)
else:
    router = None

logger = logging.getLogger(__name__)


_dsn = resolve_postgres_dsn()
if not _dsn:
    raise RuntimeError("POSTGRES_DSN is required")
try:
    _STORE = PostgresTaskStore(dsn=_dsn)
    logger.info("ingestion_task_store=postgres")
except Exception as exc:
    raise RuntimeError("failed to initialize PostgreSQL task store") from exc


def _delete_persisted_doc_rows_sync(doc_id: str, bindings: list[tuple[str, str]] | None = None) -> int:
    """删除用户绑定；仅在无剩余绑定时清理共享文档内容。"""
    if not _dsn or psycopg2 is None or psycopg2_sql is None:
        return 0
    conn = psycopg2.connect(_dsn)
    try:
        with conn:
            with conn.cursor() as cur:
                for user_id, kb_id in bindings or []:
                    cur.execute(
                        "DELETE FROM document_bindings WHERE doc_id = %s AND user_id = %s AND kb_id = %s",
                        (doc_id, user_id, kb_id),
                    )
                cur.execute("SELECT 1 FROM document_bindings WHERE doc_id = %s LIMIT 1", (doc_id,))
                if cur.fetchone() is not None:
                    return 0
                deleted_tree_nodes = 0
                for table in CLEANUP_TABLES:
                    stmt = psycopg2_sql.SQL("DELETE FROM {}.{} WHERE doc_id = %s").format(
                        psycopg2_sql.Identifier(CLEANUP_SCHEMA),
                        psycopg2_sql.Identifier(table),
                    )
                    cur.execute(stmt, (doc_id,))
                    if table == TREE_NODES_TABLE:
                        deleted_tree_nodes = int(cur.rowcount or 0)
                return deleted_tree_nodes
    finally:
        conn.close()


async def _cleanup_persisted_doc_rows(doc_id: str, bindings: list[tuple[str, str]] | None = None) -> int:
    """异步包装清理逻辑，避免阻塞事件循环。"""
    try:
        return await asyncio.to_thread(_delete_persisted_doc_rows_sync, doc_id, bindings)
    except Exception:
        logger.exception("failed to cleanup persisted doc rows: doc_id=%s", doc_id)
        return 0


_RUNNER = DocumentAssistantRunner()
_SERVICE = IngestionService(store=_STORE, runner=_RUNNER, cleanup_hook=_cleanup_persisted_doc_rows)


def create_app() -> FastAPI:
    """构建并返回 FastAPI 应用实例。"""
    if not _FASTAPI_AVAILABLE:
        raise RuntimeError("fastapi is not installed")
    # 统一默认响应类为 UTF-8，保证中文 message 在 CLI/脚本中显示稳定。
    app = FastAPI(title="Ingestion API", version="0.1.0", default_response_class=Utf8JSONResponse)
    if router is not None:
        app.include_router(router)

    @app.get(HEALTH_PATH)
    async def healthz() -> dict:
        """健康检查接口。"""
        return {"ok": True}

    @app.on_event("shutdown")
    async def archive_log_on_shutdown() -> None:
        """服务正常退出时，将本次日志归档到 OSS。"""
        try:
            from app.ingestion.log_archive import archive_service_log_to_oss

            await archive_service_log_to_oss()
        except Exception:
            logger.exception("failed to archive service log to oss")

    @app.exception_handler(RequestValidationError)
    async def validation_exception_handler(_request: Request, exc: RequestValidationError) -> JSONResponse:
        """请求参数校验异常处理器。"""
        safe_errors = jsonable_encoder(exc.errors())
        return Utf8JSONResponse(
            status_code=400,
            content={
                "code": 400,
                "message": "参数错误",
                "data": {"errors": safe_errors},
            },
        )

    @app.exception_handler(RawMineruRepairError)
    async def raw_mineru_repair_error_handler(
        _request: Request, exc: RawMineruRepairError
    ) -> JSONResponse:
        """Return the stable internal repair error contract."""
        return Utf8JSONResponse(
            status_code=exc.http_status,
            content={
                "code": exc.http_status,
                "message": str(exc),
                "data": {"error_code": exc.code, "retryable": exc.retryable},
            },
        )

    @app.exception_handler(HTTPException)
    async def http_exception_handler(_request: Request, exc: HTTPException) -> JSONResponse:
        """HTTP 异常统一处理器。"""
        code = int(exc.status_code)

        if isinstance(exc.detail, str) and exc.detail.strip():
            message = str(exc.detail)
        elif code == 400:
            message = "参数错误"
        elif code == 401:
            message = "鉴权失败"
        elif code == 404:
            message = "资源不存在"
        elif code == 409:
            message = "冲突"
        elif code == 503:
            message = "服务不可用"
        else:
            message = "服务内部错误"
        return Utf8JSONResponse(status_code=code, content={"code": code, "message": message, "data": None})

    @app.exception_handler(ServiceUnavailableError)
    async def service_unavailable_handler(_request: Request, exc: ServiceUnavailableError) -> JSONResponse:
        """服务不可用异常处理器。"""
        message = str(exc).strip() or "服务不可用"
        return Utf8JSONResponse(status_code=503, content={"code": 503, "message": message, "data": None})

    @app.exception_handler(Exception)
    async def unhandled_exception_handler(_request: Request, exc: Exception) -> JSONResponse:
        """未捕获异常统一处理器。"""
        logger.exception("unhandled exception in ingestion api", exc_info=exc)
        return Utf8JSONResponse(status_code=500, content={"code": 500, "message": "服务内部错误", "data": None})

    return app


app = create_app() if _FASTAPI_AVAILABLE else None

